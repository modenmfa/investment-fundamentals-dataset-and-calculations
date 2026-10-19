#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
spx_additions_analysis.py: S&P 500 index additions, the data and every calculation of the workbook
=====================================================================================================

LBS Investment Fundamentals, group paper, appendix code.

The strategy. Buy each stock S&P announces it will add to the S&P 500 at the open of the first trading day after
the announcement (Ann+1). Window A sells at the open of the inclusion day (the first day in the index), Window B at
its close. The benchmark is SPY bought and sold at exactly the same moments. Sample: 278 additions announced from
2010 to September 2026, on 206 announcement dates.

All the paper's numbers come from the Excel workbook "SPX Additions v<N>.xlsx", built with live formulas. This one
file reproduces it:

PART 1, DATA COLLECTION (only with --rebuild-data; it was run once, in Google Colab, with a WRDS login).
    Fills the workbook's raw data sheets from CRSP (daily stock file, via WRDS), Yahoo Finance (prices after CRSP's
    last date) and Kenneth French's data library (daily factors). It needs a WRDS account and internet access,
    so it cannot run without them; its output is the set of data sheets the calculations read.
    The candidate list of additions came from Wikipedia and from fja05680's constituent file (rebuild_addition_lists);
    each addition's announcement and inclusion dates were then taken from S&P Dow Jones Indices' press release
    ("... Set to Join S&P 500"; the links are in column W of the Clean and All sheets), by hand, giving the sheets
    'All 2010-2026 Additions' and 'Clean 2010-2026 Additions'.

PART 2, CALCULATIONS (the default). Reads ONLY the typed raw data in the workbook's data sheets:
    'Clean 2010-2026 Additions', 'Clean SPY Historical Prices', 'Event Data', 'Additions Event Window',
    'SPY Event Window', 'Additions 260-day Historicals', 'SPY 260-day Historicals', 'Fama-French Daily',
    and recomputes from scratch, with pandas, numpy and scipy, every statistic the workbook computes. It prints the
    results sheet by sheet. With --check it compares each recomputed number, date and result word with the value
    cached in the workbook, cell by cell: every number on Test 1-4, Summary, Extra Checks (v13 on) and Settings,
    and every cell of the workings sheets Trade Returns, Beta Estimation, Holding Days and Daily Account. It prints
    a match report and exits with code 1 if anything differs by more than 1e-9 (relative) + 1e-12.

How to run
----------
    python spx_additions_analysis.py                       recompute everything and print it
    python spx_additions_analysis.py --check               ... and compare with the workbook (exit 1 on a difference)
    python spx_additions_analysis.py --rebuild-data        re-collect the raw data (WRDS login asked), then recompute
    python spx_additions_analysis.py "path/to/SPX Additions v13.xlsx" --check
    python spx_additions_analysis.py --save-workings out/  also save the workings (one row per trade / day) as CSV
    python spx_additions_analysis.py --rebuild-data --rebuild-lists   also rerun the helpers that built the list
After --rebuild-data the calculations run on the new workbook version, and --check compares them with the results
cached in the original workbook (the new file has fresh data but no cached results until Excel recalculates it).

The workbook: the path given on the command line; otherwise the newest 'SPX Additions v<N>.xlsx' in this script's
folder; otherwise in ../../data/prices relative to it (the project folder layout). No file paths need editing.
Requirements: Python 3.9 or later with numpy, pandas, scipy and openpyxl. Part 1 also needs wrds and yfinance
(and requests and lxml for --rebuild-lists).
A run takes about ten seconds, most of it reading the workbook (--check reads it twice more).

What each section reproduces (workbook sheet in brackets)
---------------------------------------------------------
    Settings constants, report_settings() ........ [Settings]   same names, values and meanings (top of this file)
    rebuild_event_data() ......................... the raw data sheets                                (Part 1)
    load_inputs(), spy_close_by_date() ........... reads the raw data sheets
    beta_estimation() ............................ [Beta Estimation]   beta over days -250 to -30, minimum 100 days
    trade_returns() .............................. [Trade Returns]     stock, SPY and abnormal returns, beta-adjusted
                                                                       returns, V5, the Figure 1 CAR path
    holding_days(), daily_account() .............. [Holding Days], [Daily Account]   the strategy as an account
    newey_west_ols(), all_regressions() .......... [Regression Workings]   CAPM and four-factor alphas, Newey-West
    stat_block(), verdict() ...................... [Test 1-2 Workings], [Test 4 Workings]   t-tests, clustered t,
                                                                       Wilcoxon, sign test, MDE, Test 1's rule
    report_test1() ... report_test4() ............ [Test 1 Per Trade] [Test 2 Risk-Adjusted] [Test 3 Portfolio]
                                                   [Test 4 Robustness]
    report_summary() ............................. [Summary]
    report_extra() ............................... [Extra Checks]   E1 growth by half, E2 buying at the announcement
                                                                    close, E3 from 2020, E4 break test, E5 quarterly
    workings_reports(), check_against_workbook() . --check

Conventions copied from the spreadsheet (so the numbers agree to about 1e-12): STDEV.S divides by n - 1; RANK.AVG
gives tied values their average rank; PERCENTILE.INC is numpy's default (linear) percentile; T.DIST.2T truncates
its degrees of freedom to an integer; the Wilcoxon and Mann-Whitney z use no tie or continuity correction;
"each date once" means the average of the trades announced that day, so G = number of announcement dates.
"""
import argparse
import datetime as dt
import math
import re
import sys
from pathlib import Path
from types import SimpleNamespace

import numpy as np
import pandas as pd
from scipy import stats

# =====================================================================================================================
# SETTINGS: the same assumptions as the workbook's Settings sheet (cell in brackets)
# =====================================================================================================================
ALPHA = 0.05                                  # [C7]  significance level, two-sided
ZCRIT = stats.norm.ppf(1 - ALPHA / 2)         # [C8]  1.96: Figure 1 band and Test 3's "t > 1.96" rule
POWER = 0.80                                  # [C9]  power for the minimum detectable effect (MDE)
ZPOW = stats.norm.ppf(POWER)                  # [C10] 0.84
MDE_MULT = ZCRIT + ZPOW                       # [C11] MDE = 2.80 x clustered SE
MIN_EFFECT = 0.01                             # [C12] smallest effect worth detecting, per trade (1%)
DAYS_PER_YEAR = 252                           # [C13] annualisation
MIN_BETA_DAYS = 100                           # [C14] fewer stock returns than this: beta = 1
NW_LAGS = 5                                   # [C15] Newey-West lags (Bartlett weights)
SPLIT_YEAR = 2017                             # [C16] first half = announced in or before this year
WINSOR_LO, WINSOR_HI = 0.01, 0.99             # [C17], [C18] V3 winsorising percentiles
TRIM_N = 5                                    # [C19] V6: trades dropped at each end
EXCL_YEAR_1, EXCL_YEAR_2 = 2020, 2021         # [C20], [C21] V4: years left out
#                                               [C22] FF_LAST, the last day with Fama-French factors, is read from
#                                                     the data (31 Aug 2026): the regressions end there.
START_2020 = dt.date(2020, 1, 1)              # [C23] Extra Checks E3: trades announced on or after this day
BREAK_TRIM = 0.15                             # [C24] E4: share of dates kept out at each end of the break search
BAI_PERRON_CV_5 = 8.58                        # [C25] E4: sup-F critical value at 5% (Bai and Perron 1998, Table I:
BAI_PERRON_CV_1 = 12.29                       # [C26]     one break, one regressor, 15% trimming) and at 1%
N_DROP_V1 = 3                                 # [C27] E5: largest quarterly trades left out

# Fixed by the workbook's design (not on the Settings sheet)
FIRST_ROW = 6                                 # first event row in every per-event data sheet
BETA_FIRST_DAY, BETA_LAST_DAY = 250, 30       # beta estimated over trading days -250 to -30 before the announcement
ACCOUNT_START = dt.date(2010, 1, 1)           # Daily Account: first trading day of 2010, $1 at the 31 Dec 2009 close
HIST_DAYS = 260                               # trading days before the announcement in the 260-day sheets
WORKED_EXAMPLE = "Etsy (2020)"                # the trade shown step by step in Tests 1 and 2

# The raw data sheets the calculations read ('All 2010-2026 Additions', 'All SPY Historical Prices' and
# 'Notes & Sources' are reference only)
CLEAN, CLEAN_SPY = "Clean 2010-2026 Additions", "Clean SPY Historical Prices"
HIST, HIST_SPY = "Additions 260-day Historicals", "SPY 260-day Historicals"
WIN, WIN_SPY = "Additions Event Window", "SPY Event Window"
EVENT_DATA, FFS = "Event Data", "Fama-French Daily"

# =====================================================================================================================
# PART 1. DATA COLLECTION (runs only with --rebuild-data; needs a WRDS account and internet access)
# =====================================================================================================================
# How the data were collected. The list of additions came first: candidate additions from Wikipedia's "Selected
# changes" table and from fja05680's daily constituent file (rebuild_addition_lists below). For each addition the
# announcement and inclusion dates were then taken by hand from S&P Dow Jones Indices' press release ("... Set to Join
# S&P 500"; the links are in column W of the Clean and All sheets), giving 'All 2010-2026 Additions' (every change,
# with the dropped events in red or blue text) and 'Clean 2010-2026 Additions' (the 278 additions studied). A first
# version of the prices came from Yahoo Finance (fill_first_workbook_from_yahoo, kept for the record). Then
# rebuild_event_data() replaced every price with CRSP's, via WRDS, and built the 260-day, Event Window, Event Data
# and Fama-French sheets. It was run once, as one Google Colab cell: Colab installed wrds, openpyxl, pandas and
# yfinance, the workbook was uploaded, WRDS asked for the username, password and a Duo push, and the new workbook
# version was downloaded. Here the workbook is a plain path instead of an upload, and the result is saved next to it,
# one version up. No credentials are stored anywhere: wrds.Connection() asks for them. The builder of the analysis
# sheets later added 14 rows to 'Fama-French Daily' for the trading days of 1-21 September 2026 (dates only, with RF
# carried forward from 31 Aug, since French had not yet published them); Part 2 recreates those days itself when
# they are absent.
# Two earlier one-off Colab fixes (colab_fill_inclusion_open.py, colab_fill_missing_crsp.py) filled a few missing
# cells of v2/v3; they are superseded, because rebuild_event_data() rewrites every price cell.

LIST_START, LIST_END = "2000-01-01", "2026-09-30"
WIKI_URL = "https://en.wikipedia.org/wiki/List_of_S%26P_500_companies"
FJA_FILE = "S&P 500 Historical Components & Changes (Updated).csv"


def additions_from_wikipedia():
    """S&P 500 additions with their effective dates from Wikipedia's "Selected changes" table (announcement dates
    were then filled from S&P's press releases)."""
    from io import StringIO
    import requests
    html = requests.get(WIKI_URL, headers={"User-Agent": "Mozilla/5.0"}, timeout=30).text
    tables = pd.read_html(StringIO(html))
    changes = next(t for t in tables if isinstance(t.columns, pd.MultiIndex)
                   and any("Added" in str(c) for c in t.columns))
    changes.columns = ["_".join(str(x) for x in c if "Unnamed" not in str(x)) for c in changes.columns]
    date_col = next(c for c in changes.columns if "Date" in c)
    add_tic = next(c for c in changes.columns if c.startswith("Added") and "Ticker" in c)
    add_name = next(c for c in changes.columns if c.startswith("Added") and "Security" in c)
    rem_tic = next(c for c in changes.columns if c.startswith("Removed") and "Ticker" in c)
    reason = next((c for c in changes.columns if "Reason" in c), None)
    df = pd.DataFrame({"effective_date": pd.to_datetime(changes[date_col], errors="coerce"),
                       "ticker": changes[add_tic], "name": changes[add_name], "replaced": changes[rem_tic],
                       "reason": changes[reason] if reason else ""})
    df = df.dropna(subset=["effective_date", "ticker"])
    df = df[(df.effective_date >= LIST_START) & (df.effective_date <= LIST_END)]
    df["announce_date"] = pd.NaT                      # filled by hand from S&P DJI press releases
    df["yahoo_ticker"] = df.ticker.str.replace(".", "-", regex=False)
    return df.sort_values("effective_date").reset_index(drop=True)


def yahoo_open_close(tickers):
    """Daily open/close (split- and dividend-adjusted) from Yahoo for each ticker: the first, free price source,
    later replaced by CRSP. Yahoo has no data for many firms delisted since (the survivorship gap)."""
    import yfinance as yf
    raw = yf.download(tickers, start="1999-11-01", end="2026-10-31", auto_adjust=True, group_by="ticker",
                      progress=True, threads=True)
    rows = []
    for t in tickers:
        if t not in raw.columns.get_level_values(0):
            continue
        p = raw[t][["Open", "Close"]].dropna()
        if not p.empty:
            rows.append(p.assign(ticker=t).reset_index())
    return pd.concat(rows, ignore_index=True) if rows else pd.DataFrame()


def additions_from_fja():
    """Additions and removals from fja05680's file of daily S&P 500 constituents: a ticker present on a date but not
    on the previous row was added that day. The diff also catches ticker changes and new share classes, which were
    then classified by hand."""
    from urllib.parse import quote
    df = pd.read_csv("https://raw.githubusercontent.com/fja05680/sp500/master/" + quote(FJA_FILE))
    df["date"] = pd.to_datetime(df["date"])
    df = df.sort_values("date").reset_index(drop=True)
    members = df["tickers"].str.split(",").apply(lambda t: {x.strip() for x in t if x.strip()})
    changes = []
    for i in range(1, len(df)):
        added, removed = sorted(members[i] - members[i - 1]), sorted(members[i - 1] - members[i])
        if added or removed:
            changes.append({"date": df.loc[i, "date"].date(), "added": "|".join(added), "removed": "|".join(removed)})
    changes = pd.DataFrame(changes)
    changes = changes[(changes["date"] >= pd.Timestamp(LIST_START).date())
                      & (changes["date"] <= pd.Timestamp(LIST_END).date())]
    additions = (changes[changes["added"] != ""].assign(ticker=lambda d: d["added"].str.split("|"))
                 .explode("ticker")[["date", "ticker", "removed"]]
                 .rename(columns={"date": "first_day_in_index", "removed": "removed_same_day"}))
    return changes, additions


def rebuild_addition_lists(folder):
    """Rerun the two list builders and save their output as CSV files in folder (--rebuild-data --rebuild-lists)."""
    folder = Path(folder)
    adds = additions_from_wikipedia()
    adds.to_csv(folder / "additions_raw.csv", index=False)
    print(f"{len(adds)} additions between {LIST_START} and {LIST_END} (Wikipedia)")
    prices = yahoo_open_close(sorted(set(adds.yahoo_ticker)) + ["^GSPC", "SPY"])
    prices.to_csv(folder / "prices.csv", index=False)
    have = set(prices.ticker) if len(prices) else set()
    adds[~adds.yahoo_ticker.isin(have)].to_csv(folder / "missing.csv", index=False)
    print(f"Yahoo has prices for {adds.yahoo_ticker.isin(have).sum()} of {len(adds)} additions. Caution: Yahoo "
          f"tickers are today's symbols, so a reused ticker can return a different company.")
    changes, additions = additions_from_fja()
    changes.to_csv(folder / "fja_changes_2000_2026.csv", index=False)
    additions.to_csv(folder / "fja_additions_2000_2026.csv", index=False)
    print(f"{len(changes)} change dates, {len(additions)} added tickers (fja05680)")


def fill_first_workbook_from_yahoo(src):
    """The first price fill (kept for the record and not called by main(): rebuild_event_data rewrites every price).
    Sheet 1 rows 6+: the price for each date column (F, I, L, O, R -> G, J, M, P, S; Close or Open as the row-5
    header says) and the volume / 1e6 (H, K, N, Q, T) from Yahoo, split-adjusted but not dividend-adjusted; a note
    in column Z for every value that could not be found. Sheet 2: the S&P 500 index (^GSPC) on the same dates.
    A ticker Yahoo now lists under another company's name is not filled (reused ticker)."""
    import time
    import yfinance as yf
    from openpyxl import load_workbook
    src = Path(src)
    out = src.with_name(src.stem + " - filled.xlsx")
    date_cols = {"F": ("G", "H"), "I": ("J", "K"), "L": ("M", "N"), "O": ("P", "Q"), "R": ("S", "T")}
    alias = {"BRK.B": "BRK-B", "FB": "META", "NLOK": "GEN", "PEAK": "DOC", "DWDP": "DD", "HFC": "DINO",
             "WLTW": "WTW", "RE": "EG", "CDAY": "DAY", "FLT": "CPAY", "FBHS": "FBIN", "ADS": "BFH", "KORS": "CPRI"}
    generic = {"inc", "corp", "corporation", "co", "company", "ltd", "plc", "the", "group", "holdings", "holding",
               "class", "a", "b", "c", "and", "of", "nv", "sa", "lp", "trust", "international", "technologies",
               "technology", "systems", "industries", "s", "communications", "new"}
    words = lambda s: {w for w in re.findall(r"[a-z0-9]+", str(s).lower()) if w not in generic and len(w) > 1}

    def history(t, start, end):
        for attempt in range(3):
            try:
                h = yf.Ticker(t).history(start=start, end=end, auto_adjust=False, actions=False)
                if h is not None and not h.empty:
                    h.index = pd.to_datetime(h.index).tz_localize(None).normalize()
                    return h
                return None
            except Exception:
                time.sleep(2 * (attempt + 1))
        return None

    def yahoo_name(t):
        try:
            i = yf.Ticker(t).get_info()
            return i.get("longName") or i.get("shortName") or ""
        except Exception:
            return ""

    def as_date(v):
        try:
            return pd.Timestamp(v).normalize() if v not in (None, "") else None
        except Exception:                               # an unreadable date cell is skipped
            return None

    wb = load_workbook(src)
    ws, spx = wb.worksheets[0], wb.worksheets[1]
    kind = {c: str(ws[f"{c}5"].value or "").lower() for c, _ in date_cols.values()}
    rows = [r for r in range(6, ws.max_row + 1) if ws[f"D{r}"].value and ws[f"F{r}"].value]
    all_dates = [d for d in (as_date(ws[f"{c}{r}"].value) for r in rows for c in date_cols) if d is not None]
    lo, hi = min(all_dates) - pd.Timedelta(days=10), max(all_dates) + pd.Timedelta(days=10)
    idx = history("^GSPC", lo, hi)
    for r in rows:
        for dcol, (pcol, _) in date_cols.items():
            d = as_date(ws[f"{dcol}{r}"].value)
            if idx is not None and d in idx.index:
                spx[f"{pcol}{r}"] = float(idx.loc[d, "Open" if "open" in kind[pcol] else "Close"])
                spx[f"{pcol}{r}"].number_format = "0.00"
    cache, log = {}, []
    for r in rows:
        name, tick = ws[f"C{r}"].value, str(ws[f"D{r}"].value).strip()
        ytick = alias.get(tick, tick.replace(".", "-"))
        if ytick not in cache:
            h = history(ytick, lo, hi)
            cache[ytick] = (h, yahoo_name(ytick) if h is not None else "")
        h, yname = cache[ytick]
        notes = []
        if h is None:
            notes.append(f"No Yahoo data for {ytick} (likely delisted/acquired; Yahoo drops delisted tickers)")
        elif tick not in alias and yname and not (words(name) & words(yname)):
            notes.append(f"Not filled: Yahoo lists {ytick} as '{yname}', not '{name}' (ticker reused or renamed)")
            h = None
        if h is not None:
            for dcol, (pcol, vcol) in date_cols.items():
                d = as_date(ws[f"{dcol}{r}"].value)
                if d is None:
                    continue
                if d not in h.index:
                    notes.append(f"No Yahoo row for {ytick} on {d.date()}")
                    continue
                ws[f"{pcol}{r}"] = float(h.loc[d, "Open" if "open" in kind[pcol] else "Close"])
                ws[f"{pcol}{r}"].number_format = "0.00"
                ws[f"{vcol}{r}"] = float(h.loc[d, "Volume"]) / 1e6
                ws[f"{vcol}{r}"].number_format = "0"
        if ytick != tick.replace(".", "-") and h is not None:
            notes.insert(0, f"Yahoo ticker used: {ytick}")
        if notes:
            old = ws[f"Z{r}"].value
            ws[f"Z{r}"] = ((str(old) + " | ") if old else "") + "; ".join(notes)
        log.append((r, tick, ytick, yname, "; ".join(notes) or "ok"))
    if not ws["Z5"].value:
        ws["Z5"] = "Data notes"
    wb.save(out)
    pd.DataFrame(log, columns=["row", "ticker", "yahoo_ticker", "yahoo_name", "result"]).to_csv(
        out.with_suffix(".log.csv"), index=False)
    return out


def rebuild_event_data(src_path):
    """Fill the SPX Additions workbook with CRSP prices (via WRDS), SPY, and the Fama-French factors.
    (colab_build_event_data.py, carried over unchanged apart from the Colab upload/download.)

    CRSP is the source wherever it has data. Dates after CRSP's last date (SPY, and additions whose window ends after
    it) come from Yahoo Finance, and every row says which source it came from.
    Input : the workbook; the additions to use are read from "Clean 2010-2026 Additions".
    Output: the same workbook, saved one version up (v13 -> v14), with
      - the price and volume cells of the Clean and All sheets (stock and SPY) replaced by CRSP / Yahoo values;
        red and blue rows are left alone
      - "Additions 260-day Historicals" and "SPY 260-day Historicals" filled: block k = k trading days before the
        announcement (date, open, close, volume in millions)
      - sheets in the same layout: Event Data (one row per addition: source, CRSP match, size, liquidity,
        quarterly flag, checks), Additions Event Window (every trading day from the announcement, block 0, to the
        inclusion day), SPY Event Window (SPY on the same days), Fama-French Daily (Mkt-RF, SMB, HML, RF, Mom),
        Needs Review (additions whose match or prices need a look) and Notes & Sources.
    Prices are adjusted for splits to the share basis of the announcement day, so a split inside the 260 days does
    not show up as a price jump. On and after the announcement day they equal CRSP's own prices.
    WRDS asks for the username and password (and a Duo push). Takes several minutes.
    """
    try:
        import io, time, zipfile, urllib.request                                       # noqa: E401
        from openpyxl import load_workbook
        from openpyxl.styles import Font, Border, Side, Alignment
        from openpyxl.utils import get_column_letter as L
        import wrds                                                                    # noqa: F401
        import yfinance                                                                # noqa: F401
    except ImportError as err:
        sys.exit(f"Part 1 needs the wrds and yfinance packages ({err}). In Colab: pip install wrds yfinance openpyxl")

    # ---------------------------------------------------------------- settings
    # In Colab the workbook was uploaded (google.colab files.upload); here it is the path passed in. The output is
    # written next to it, one version up (v13 -> v14).
    SRC = str(src_path)
    m = re.search(r"v(\d+)", Path(SRC).name)
    OUT = str(Path(SRC).with_name(Path(SRC).name[:m.start()] + f"v{int(m.group(1)) + 1}" + Path(SRC).name[m.end():])) if m \
        else SRC.replace(".xlsx", " - filled.xlsx")
    CLEAN, CLEAN_SPY = "Clean 2010-2026 Additions", "Clean SPY Historical Prices"
    ALL, ALL_SPY = "All 2010-2026 Additions", "All SPY Historical Prices"
    HIST, HIST_SPY = "Additions 260-day Historicals", "SPY 260-day Historicals"
    WIN, WIN_SPY = "Additions Event Window", "SPY Event Window"
    EVD, FFS, REV, NOTES = "Event Data", "Fama-French Daily", "Needs Review", "Notes & Sources"
    HIST_DAYS = 260                       # trading days before the announcement in the 260-day sheets
    PRE_DAYS, POST_DAYS = 260, 5          # trading days pulled before the announcement / after inclusion (for the Clean sheets' Inclusion +1 and +3 columns)
    SPY_PERMNO = 84398
    CHECK_TOL = 0.01                      # flag if our window return and the workbook's old one differ by > 1 point
    OVERRIDES = {}                        # {"Company name exactly as in the sheet": permno} to force a match after review
    FRENCH = "https://mba.tuck.dartmouth.edu/pages/faculty/ken.french/ftp/"

    # ---------------------------------------------------------------- 1. events from the Clean sheet
    COLS = {2: "name", 3: "ticker", 4: "sector", 5: "ann", 6: "wb_ann_c", 8: "a1", 9: "wb_a1_o",
            11: "inc", 12: "wb_inc_o", 13: "wb_inc_c", 15: "i1", 16: "wb_i1_c", 18: "i3", 19: "wb_i3_c"}
    ws = load_workbook(SRC, data_only=True)[CLEAN]
    recs = []
    for r, v in enumerate(ws.iter_rows(min_row=6, values_only=True), start=6):
        if v[2] is not None and v[1] is None:                       # data rows only (skips the key at the bottom)
            recs.append(dict({k: v[i] for i, k in COLS.items()}, row=r))
    ev = pd.DataFrame(recs)
    for c in ["ann", "a1", "inc", "i1", "i3"]:
        ev[c] = pd.to_datetime(ev[c], errors="coerce")
    bad = ev[ev[["ann", "a1", "inc", "i1", "i3"]].isna().any(axis=1)]
    if len(bad):
        print("Skipped (a date is missing in the Clean sheet):", ", ".join(bad["name"].astype(str)))
        ev = ev.drop(bad.index)
    for c in [c for c in ev if c.startswith("wb_")]:
        ev[c] = pd.to_numeric(ev[c], errors="coerce")
    ev["ticker"] = ev.ticker.astype(str).str.strip().str.upper()
    ev = ev.sort_values(["ann", "row"]).reset_index(drop=True)
    ev.insert(0, "event_id", [f"E{i + 1:03d}" for i in range(len(ev))])
    print(f"{len(ev)} events, announced {ev.ann.min().date()} to {ev.ann.max().date()}")

    # quarterly rebalance: effective on the Monday after the third Friday of Mar/Jun/Sep/Dec (or the next
    # trading day), announced at least 7 calendar days earlier. Additions effective that day at short notice
    # are replacements that happen to share the date, so they are not counted as scheduled.
    def third_friday(y, m):
        d = dt.date(y, m, 15)
        return d + dt.timedelta((4 - d.weekday()) % 7)
    ev["on_rebalance_date"] = [d.month in (3, 6, 9, 12) and 1 <= (d.date() - third_friday(d.year, d.month)).days <= 4
                               for d in ev.inc]
    ev["gap_days"] = (ev.inc - ev.ann).dt.days
    ev["quarterly"] = ev.on_rebalance_date & (ev.gap_days >= 7)
    # Checked against S&P's press releases (data/prices/quarterly_rebalance_check.csv). These take effect on a
    # rebalance date but replace a company that left because of a takeover, spinoff or conversion, so they are
    # not part of the scheduled quarterly batch.
    OFF_CYCLE = {
        ("CVC", "2010-12-09"): "replaced King Pharmaceuticals (acquired by Pfizer)",
        ("ZTS", "2013-06-14"): "added because Pfizer was splitting off its Zoetis stake",
        ("AAL", "2015-03-16"): "replaced Allergan (acquired by Actavis)",
        ("HLT", "2017-06-09"): "replaced Yahoo! (converting to a closed-end fund)",
        ("RE", "2017-06-12"): "replaced Mead Johnson (acquired by Reckitt Benckiser)",
        ("CDNS", "2017-09-08"): "replaced Staples (acquired by Sycamore Partners)",
        ("FSLR", "2022-12-12"): "replaced Fortune Brands (moved to MidCap 400 after its MasterBrand spinoff)",
    }
    ev["quarterly_note"] = [OFF_CYCLE.get((t, d.strftime("%Y-%m-%d")), "") for t, d in zip(ev.ticker, ev.ann)]
    ev.loc[ev.quarterly_note != "", "quarterly"] = False
    ev["quarterly_note"] = ["Off-cycle per S&P's release: " + x if x else "" for x in ev.quarterly_note]

    # ---------------------------------------------------------------- 2. Fama-French daily factors (no login)
    def french(zipname):
        req = urllib.request.Request(FRENCH + zipname, headers={"User-Agent": "Mozilla/5.0"})
        raw = urllib.request.urlopen(req, timeout=60).read()
        z = zipfile.ZipFile(io.BytesIO(raw))
        lines = z.read(z.namelist()[0]).decode("latin-1").splitlines()
        header, rows = None, []
        for line in lines:
            parts = [p.strip() for p in line.split(",")]
            if re.fullmatch(r"\d{8}", parts[0]):
                rows.append(parts)
            elif rows:
                break                                    # first table only (daily); stop at the end of it
            elif len(parts) > 1:
                header = parts[1:]                       # the last comma line before the data holds the names
        df = pd.DataFrame(rows, columns=["date"] + header)
        df["date"] = pd.to_datetime(df.date, format="%Y%m%d")
        return df.set_index("date").apply(pd.to_numeric, errors="coerce") / 100   # percent -> decimal
    ff = french("F-F_Research_Data_Factors_daily_CSV.zip").join(french("F-F_Momentum_Factor_daily_CSV.zip"), how="left")
    ff.columns = [c.replace(" ", "") for c in ff.columns]
    ff = ff[ff.index >= "2008-01-01"]
    print(f"Fama-French daily: {ff.index.min().date()} to {ff.index.max().date()}, columns {list(ff.columns)}")

    # ---------------------------------------------------------------- 3. CRSP via WRDS
    import wrds
    db = wrds.Connection()                               # asks for your WRDS username + password (and Duo)

    def columns(lib, table):
        try:
            return set(db.describe_table(lib, table)["name"].str.lower())
        except Exception:
            return None

    def pick(have, *cands):
        return next((c for c in cands if c in have), None)

    # CRSP has two formats: the legacy daily file (crsp.dsf) and the newer one (crsp.dsf_v2). The legacy file
    # stops at the end of 2024, so later dates come from the new file. Column names are looked up, not assumed.
    PRICE_TABLES = [
        ("dsf", dict(date=("date",), open=("openprc",), close=("prc",), ret=("ret",), retx=("retx",),
                     vol=("vol",), shrout=("shrout",))),
        ("dsf_v2", dict(date=("dlycaldt", "date"), open=("dlyopen", "openprc"), close=("dlyprc", "prc"),
                        ret=("dlyret", "ret"), retx=("dlyretx", "retx"), vol=("dlyvol", "vol"),
                        shrout=("shrout", "dlyshrout"))),
    ]
    sources = []
    for table, cands in PRICE_TABLES:
        have = columns("crsp", table)
        if not have:
            print(f"crsp.{table}: not available"); continue
        sel = {k: pick(have, *v) for k, v in cands.items()}
        if not (sel["date"] and sel["close"] and sel["ret"]):
            print(f"crsp.{table}: columns not recognised, skipped. Columns: {sorted(have)}"); continue
        if not sel["retx"]:
            print(f"crsp.{table}: no price-return column, so splits cannot be detected in its rows")
        last = db.raw_sql(f"select max({sel['date']}) as d from crsp.{table} where permno = {SPY_PERMNO}").d.iloc[0]
        sources.append(dict(table=table, sel=sel, last=pd.Timestamp(last)))
        print(f"crsp.{table}: data to {pd.Timestamp(last).date()}")
    if not sources:
        raise SystemExit("No CRSP daily table found. Check the table names printed above.")

    def pull(src, where):
        sel = src["sel"]
        fields = ", ".join(f"{c} as {k}" for k, c in sel.items() if c)
        df = db.raw_sql(f"select permno, {fields} from crsp.{src['table']} where {where}", date_cols=["date"])
        for k in ["open", "close", "ret", "retx", "vol", "shrout"]:
            if k not in df:
                df[k] = np.nan
        df["src"] = src["table"]
        return df

    def pull_all(where_for):           # where_for(src) -> SQL condition; later tables only fill dates after earlier ones
        out, covered = [], pd.Timestamp("1900-01-01")
        for s in sources:
            datecol = s["sel"]["date"]
            if s["last"] <= covered:
                continue
            df = pull(s, f"({where_for(s)}) and {datecol} > '{covered.date()}'")
            out.append(df); covered = s["last"]
        return pd.concat(out, ignore_index=True)

    def is_power_of_10(x):
        return pd.notna(x) and x > 0 and abs(np.log10(x) - round(np.log10(x))) < 0.01 and round(np.log10(x)) != 0

    # SPY over the whole sample, and a units check between the two CRSP formats on overlapping dates
    spy = pull_all(lambda s: f"permno = {SPY_PERMNO} and {s['sel']['date']} >= '2008-01-01'")
    scale = {}
    if len(sources) == 2:
        a = pull(sources[0], f"permno = {SPY_PERMNO} and {sources[0]['sel']['date']} >= '2023-01-01'").set_index("date")
        b = pull(sources[1], f"permno = {SPY_PERMNO} and {sources[1]['sel']['date']} >= '2023-01-01'").set_index("date")
        both = a.index.intersection(b.index)
        for k in ["close", "open", "ret", "vol", "shrout"]:
            r = (a.loc[both, k].abs() / b.loc[both, k].abs()).replace([np.inf, -np.inf], np.nan).median()
            scale[k] = float(r) if pd.notna(r) else np.nan
        print("Legacy / new CRSP ratios on overlapping SPY dates (should be 1):", {k: round(v, 4) for k, v in scale.items()})
        for k in ["vol", "shrout"]:                      # unit differences (shares vs thousands) are fixed here
            if is_power_of_10(scale.get(k)):
                spy.loc[spy.src == sources[1]["table"], k] *= scale[k]
                print(f"  {k}: new-format values rescaled by {scale[k]:g} to match the legacy units")
    spy = spy.drop_duplicates("date").sort_values("date").set_index("date")
    spy["close"] = spy.close.abs()
    crsp_last = max(s["last"] for s in sources)

    # ---------------------------------------------------------------- 3b. Yahoo Finance after CRSP's last date
    import yfinance as yf
    def yahoo(ticker, start, end):
        """Daily open/close/volume and returns from Yahoo. Close is split-adjusted only; ret uses Adj Close
        (dividends included), retx uses Close. Starts 15 days early so the first return can be computed."""
        for attempt in range(3):
            try:
                h = yf.Ticker(ticker).history(start=start - pd.Timedelta(days=15), end=end + pd.Timedelta(days=1),
                                              auto_adjust=False, actions=True)
                break
            except Exception as err:
                print(f"  Yahoo {ticker}: {err}; retrying"); time.sleep(10)
        else:
            return pd.DataFrame()
        if h.empty:
            return pd.DataFrame()
        h.index = pd.to_datetime(h.index).tz_localize(None).normalize()
        out = pd.DataFrame({"open": h["Open"], "close": h["Close"], "vol": h["Volume"],
                            "ret": h["Adj Close"].pct_change(), "retx": h["Close"].pct_change()})
        out.index.name = "date"
        return out[out.index >= start]

    ys = yahoo("SPY", crsp_last - pd.Timedelta(days=40), pd.Timestamp.today().normalize())
    seam = ys.index.intersection(spy.index)
    if len(seam):
        gap = (ys.loc[seam, "close"] / spy.loc[seam, "close"] - 1).abs().max()
        print(f"SPY: Yahoo vs CRSP close on {len(seam)} overlapping days, largest difference {gap:.3%}")
    ys = ys[ys.index > crsp_last].assign(src="yahoo")
    spy = pd.concat([spy, ys]).sort_index()
    print(f"SPY: CRSP to {crsp_last.date()}, Yahoo to {spy.index.max().date()}")
    cal = spy.index                                     # trading-day calendar = days SPY traded

    # ---------------------------------------------------------------- 4. ticker -> PERMNO on the announcement date
    GENERIC = {"inc", "corp", "corporation", "co", "company", "ltd", "plc", "the", "group", "holdings", "holding",
               "class", "and", "of", "nv", "sa", "lp", "trust", "international", "technologies", "technology",
               "systems", "industries", "communications", "new", "com", "hldgs", "intl", "cos", "del", "a", "b"}
    def words(s):                                    # CRSP spells initials apart ("P V H CORP"): join them first
        s = re.sub(r"\s+", " ", str(s).lower().replace("&", ""))
        s = re.sub(r"\b(\w) (?=\w\b)", r"\1", s)
        return {w for w in re.findall(r"[a-z0-9]+", s) if w not in GENERIC and len(w) > 1}
    score = lambda other_name, our_name: len(words(other_name) & words(our_name))

    NAME_TABLES = [("stocknames", dict(ticker=("ticker",), comnam=("comnam",), shrcls=("shrcls",),
                                       start=("namedt",), end=("nameenddt",))),
                   ("stocknames_v2", dict(ticker=("ticker",), comnam=("comnam", "issuernm", "securitynm"),
                                          shrcls=("shrcls", "shareclass"), start=("namedt", "secinfostartdt"),
                                          end=("nameenddt", "secinfoenddt")))]
    name_srcs = []
    for table, cands in NAME_TABLES:
        have = columns("crsp", table)
        if not have:
            continue
        sel = {k: pick(have, *v) for k, v in cands.items()}
        if all(sel[k] for k in ["ticker", "comnam", "start", "end"]):
            name_srcs.append((table, sel))
    print("Name tables used:", [t for t, _ in name_srcs])
    if not name_srcs:
        raise SystemExit("No CRSP name table found. Check the table names printed above.")

    def names_where(cond_for):
        out = []
        for table, sel in name_srcs:
            f = ", ".join(f"{c} as {k}" for k, c in sel.items() if c)
            df = db.raw_sql(f"select permno, {f} from crsp.{table} where {cond_for(sel)}", date_cols=["start", "end"])
            if "shrcls" not in df:
                df["shrcls"] = None
            out.append(df)
        df = pd.concat(out, ignore_index=True) if out else pd.DataFrame()
        if len(df):
            df["end"] = df["end"].fillna(pd.Timestamp("2100-01-01"))
        return df

    ev["beyond_crsp"] = ev.i3 > crsp_last                  # window ends after CRSP's last date: whole event from Yahoo
    ev["price_source"] = np.where(ev.beyond_crsp, "Yahoo", "CRSP")
    variants = sorted({v for t in ev.ticker for v in (t, t.replace(".", ""), t.split(".")[0])})
    names = names_where(lambda sel: f"{sel['ticker']} in ({','.join(repr(v) for v in variants)})")

    def match(e):
        if e["beyond_crsp"]:
            return None, None, None, "Yahoo by ticker (after CRSP's last date)"
        if e["name"] in OVERRIDES:
            p = OVERRIDES[e["name"]]
            n = names[names.permno == p]
            return p, (n.comnam.iloc[-1] if len(n) else ""), "", "manual override"
        t = e["ticker"]; base, _, cls = t.partition(".")
        cand = names[names.ticker.isin({t, t.replace(".", ""), base})]
        for d, how in [(e["ann"], "ticker on announcement date"), (e["inc"], "ticker on inclusion date")]:
            at = cand[(cand["start"] <= d) & (cand["end"] >= d)]
            if cls and (at.shrcls == cls).any():
                at = at[at.shrcls == cls]
            if len(at):
                at = at.assign(s=at.comnam.map(lambda x: score(x, e["name"]))).sort_values("s")
                b = at.iloc[-1]
                return int(b.permno), b.comnam, b.ticker, how
        # ticker not valid on either date: search CRSP names that were active on the inclusion date
        w = sorted(words(e["name"]), key=len, reverse=True)
        if not w:
            return None, None, None, "not found (no usable name)"
        d = e["inc"].date()
        hits = names_where(lambda sel: f"{sel['comnam']} ilike '%%{w[0].upper()}%%' and "
                                       f"{sel['start']} <= '{d}' and coalesce({sel['end']}, '2100-01-01') >= '{d}'")
        if hits.empty:
            return None, None, None, f"not found (no CRSP name containing '{w[0]}' on {d})"
        hits = hits.assign(s=hits.comnam.map(lambda x: score(x, e["name"]))).sort_values("s")
        b = hits.iloc[-1]
        return int(b.permno), b.comnam, b.ticker, "company-name search on inclusion date"

    mt = pd.DataFrame([match(e) for _, e in ev.iterrows()], columns=["permno", "src_name", "src_ticker", "match_method"])
    ev = pd.concat([ev, mt], axis=1)
    print(ev.match_method.value_counts().to_string())

    # ---------------------------------------------------------------- 5. daily stock data per event
    def window(e):
        i0 = cal.searchsorted(e["ann"])                         # announcement day (or next trading day)
        i1 = min(cal.searchsorted(e["inc"]) + POST_DAYS, len(cal) - 1)
        return cal[max(i0 - PRE_DAYS, 0)], cal[i1]

    ok = ev[ev.permno.notna() & ~ev.beyond_crsp]
    raw = []
    for i in range(0, len(ok), 60):
        chunk = ok.iloc[i:i + 60]
        def cond(s, chunk=chunk):
            dc = s["sel"]["date"]
            return " or ".join(f"(permno = {int(e.permno)} and {dc} between '{a.date()}' and '{b.date()}')"
                               for _, e in chunk.iterrows() for a, b in [window(e)])
        raw.append(pull_all(cond))
    raw = pd.concat(raw, ignore_index=True).drop_duplicates(["permno", "date"]) if raw else pd.DataFrame()
    if len(raw):
        for k in ["vol", "shrout"]:
            if is_power_of_10(scale.get(k)):
                raw.loc[raw.src == sources[1]["table"], k] *= scale[k]
        raw["quote_is_midpoint"] = raw.close < 0            # CRSP: negative price = bid/ask midpoint, no trade
        raw["close"] = raw.close.abs()

    panel = []
    for _, e in ok.iterrows():
        a, b = window(e)
        d = raw[(raw.permno == e.permno) & raw.date.between(a, b)].copy()
        d["event_id"] = e.event_id
        panel.append(d)

    # share counts entered by hand where Yahoo has none: listed class only, the way CRSP counts shares
    SHARES_BY_HAND = {
        ("ECHO", "2026-03-06"): 159_142_240,   # EchoStar Class A, 10-Q cover page as of 2026-07-21; Class B is unlisted
    }
    for _, e in ev[ev.beyond_crsp].iterrows():
        a, b = window(e)
        yt = e["ticker"].replace(".", "-")
        h = yahoo(yt, a, b)
        if not len(h):
            print(f"  Yahoo: no data for {yt} ({e['name']})"); continue
        h = h[(h.index >= a) & (h.index <= b)].reset_index()
        shares = np.nan                                  # shares outstanding near the announcement, in thousands
        try:
            sh = yf.Ticker(yt).get_shares_full(start=e["ann"] - pd.Timedelta(days=120), end=e["ann"] + pd.Timedelta(days=1))
            if sh is not None and len(sh):
                shares = float(sh.iloc[-1]) / 1000
        except Exception:
            pass
        by_hand = SHARES_BY_HAND.get((e["ticker"], e["ann"].strftime("%Y-%m-%d")))
        if by_hand:
            shares = by_hand / 1000
        try:                                             # Yahoo's company name, to check the ticker is the right firm
            yname = yf.Ticker(yt).info.get("longName") or ""
        except Exception:
            yname = ""
        ev.loc[ev.event_id == e.event_id, "src_name"] = yname
        panel.append(h.assign(event_id=e.event_id, permno=np.nan, shrout=shares, quote_is_midpoint=False, src="yahoo"))

    daily = pd.concat(panel, ignore_index=True)
    for c in ["open", "close", "ret", "retx", "vol", "shrout", "permno"]:
        daily[c] = pd.to_numeric(daily[c], errors="coerce")
    daily["date"] = pd.to_datetime(daily["date"])
    ev["name_score"] = [score(n, o) if isinstance(n, str) and n else 0 for n, o in zip(ev["src_name"], ev["name"])]

    # rel_day: trading days from the announcement day (0). Split adjustment: a split shows up as a day where the
    # price ratio and CRSP's price return disagree; prices before it are put on the announcement day's share basis.
    ann_of = ev.set_index("event_id")["ann"]
    adj = []
    for eid, d in daily.groupby("event_id"):
        d = d.sort_values("date").copy()
        d["rel_day"] = cal.searchsorted(d.date) - cal.searchsorted(ann_of[eid])
        ratio = (d.close / d.close.shift(1)) / (1 + d.retx)
        ratio = ratio.where(np.log(ratio.abs()).abs() > 0.01, 1.0).fillna(1.0)
        ratio[d.src == "yahoo"] = 1.0                    # Yahoo prices are already split-adjusted
        f = ratio.cumprod()
        base = f[d.rel_day <= 0].iloc[-1] if (d.rel_day <= 0).any() else f.iloc[0]
        k = f / base
        d["open"], d["close"], d["vol"] = d.open / k, d.close / k, d.vol * k
        d["split_ratio"] = ratio
        adj.append(d)
    daily = pd.concat(adj, ignore_index=True)
    print(f"Splits adjusted: {int((daily.split_ratio != 1).sum())} split days across "
          f"{daily[daily.split_ratio != 1].event_id.nunique()} additions")

    # ---------------------------------------------------------------- 6. event-level fields
    def on(df, day, col):
        r = df[df.date == day]
        return float(r[col].iloc[0]) if len(r) and pd.notna(r[col].iloc[0]) else np.nan

    rows = []
    for _, e in ev.iterrows():
        d = daily[daily.event_id == e.event_id]
        pre_inc = cal[cal.searchsorted(e["inc"]) - 1]
        r = dict(event_id=e.event_id, pre_inc=pre_inc)
        for lab, day, col in [("ann_c", e.ann, "close"), ("a1_o", e.a1, "open"), ("pre_inc_c", pre_inc, "close"),
                              ("inc_o", e.inc, "open"), ("inc_c", e.inc, "close"), ("i1_c", e.i1, "close"),
                              ("i3_c", e.i3, "close")]:
            r[lab] = on(d, day, col)
            r["spy_" + lab] = float(spy.loc[day, col]) if day in spy.index else np.nan
        for lab, day in [("ann_v", e.ann), ("a1_v", e.a1), ("inc_v", e.inc), ("i1_v", e.i1), ("i3_v", e.i3)]:
            r[lab] = on(d, day, "vol")
            r["spy_" + lab] = float(spy.loc[day, "vol"]) if day in spy.index else np.nan
        r["mktcap_ann_musd"] = r["ann_c"] * on(d, e.ann, "shrout") / 1000        # shrout is in thousands
        est = d[d.rel_day.between(-60, -11)]
        r["adv_dollar_musd"] = (est.vol * est.close).mean() / 1e6 if len(est) else np.nan
        r["n_est_days"] = int(d.rel_day.between(-250, -30).sum())                 # days available for beta
        inwin = d[(d.date > e.a1) & (d.date <= e.inc)]
        r["div_yield_window"] = float((inwin.ret - inwin.retx).fillna(0).sum()) if len(inwin) else np.nan
        span = d[(d.date >= e.ann) & (d.date <= e.i3)]
        r["split_in_window"] = bool((span.split_ratio != 1).any())
        r["midpoint_in_window"] = bool(span.quote_is_midpoint.fillna(False).any())
        rows.append(r)
    ev = ev.merge(pd.DataFrame(rows), on="event_id", how="left")
    ev["demand_days"] = ev.mktcap_ann_musd / ev.adv_dollar_musd     # market cap / normal dollar volume

    # checks against the prices that were in the workbook before (window returns, so later splits don't matter)
    ev["chk_A_diff"] = (ev.inc_o / ev.a1_o) - (ev.wb_inc_o / ev.wb_a1_o)
    ev["chk_B_diff"] = (ev.inc_c / ev.a1_o) - (ev.wb_inc_c / ev.wb_a1_o)

    def problems(e):
        p = []
        if e.price_source == "CRSP" and pd.isna(e.permno):
            p.append(e.match_method)
        elif e.price_source == "Yahoo" and not e.src_name:
            p.append("Yahoo returned no company name: check the ticker")
        elif e.name_score == 0:
            p.append(f"name check: {e.price_source} name '{e.src_name}' shares no word with '{e['name']}'")
        if e.price_source == "Yahoo" or pd.notna(e.permno):
            for lab, txt in [("ann_c", "announcement close"), ("a1_o", "Ann+1 open"), ("inc_o", "inclusion open"),
                             ("inc_c", "inclusion close"), ("i1_c", "Inc+1 close"), ("i3_c", "Inc+3 close")]:
                if pd.isna(e[lab]): p.append(f"no {e.price_source} {txt}")
            for w in ["A", "B"]:
                x = e[f"chk_{w}_diff"]
                if pd.notna(x) and abs(x) > CHECK_TOL: p.append(f"window {w} return differs from the old workbook value by {x:+.1%}")
            if e.split_in_window: p.append("split inside the window (prices adjusted)")
            if e.midpoint_in_window: p.append("a close in the window is a bid/ask midpoint")
            if pd.isna(e.mktcap_ann_musd): p.append("no share count: market cap missing")
        return "; ".join(p)
    ev["review"] = ev.apply(problems, axis=1)
    print(f"\n{(ev.review != '').sum()} additions need a look (sheet '{REV}').")
    print(f"Additions with full window prices: {ev[['a1_o', 'inc_o', 'inc_c']].notna().all(axis=1).sum()} of {len(ev)}")

    # ---------------------------------------------------------------- 7. write into the workbook, in its own style
    book = load_workbook(SRC)                       # formulas and formatting kept
    for sn in [EVD, WIN, WIN_SPY, FFS, REV, NOTES]:  # sheets this script makes are rebuilt on every run
        if sn in book.sheetnames:
            del book[sn]

    THIN, MED = Side(style="thin"), Side(style="medium")
    F_TITLE, F_BOLD, F_NORM = Font(name="Calibri", size=14, bold=True), Font(name="Calibri", size=11, bold=True), Font(name="Calibri", size=11)
    A_C, A_CC, A_L = Alignment(horizontal="center", vertical="center"), Alignment(horizontal="centerContinuous", vertical="center"), Alignment(horizontal="left")
    B_L, B_LB, B_B = Border(left=THIN), Border(left=THIN, bottom=THIN), Border(bottom=THIN)
    B_BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
    FMT = {"date": "yyyy-mm-dd", "num": "0.00", "int": "0", "pct": "0.00%", "pct3": "0.000%", "text": "General", "bold": "General"}

    def title(ws, text, last_col):
        ws.sheet_view.showGridLines = False
        ws.sheet_view.zoomScale = 77
        ws.column_dimensions["A"].width = 2.7
        ws["B2"] = text; ws["B2"].font = F_TITLE
        for c in range(2, last_col + 1):
            ws.cell(2, c).border = Border(bottom=MED)
        ws.row_dimensions[2].height = 18.75
        ws.row_dimensions[3].height = 5.1

    def clean(v):
        if v is None or v is pd.NaT or (isinstance(v, (float, np.floating)) and np.isnan(v)):
            return None
        if isinstance(v, np.bool_): return bool(v)
        if isinstance(v, (np.floating, np.integer)): return v.item()
        if isinstance(v, pd.Timestamp): return v.to_pydatetime()
        return v

    def put(ws, r, c, v, kind="text", left=False, bottom=False, right=False):
        cell = ws.cell(r, c)
        cell.value = clean(v)
        cell.number_format = FMT[kind]
        cell.font = F_BOLD if kind == "bold" else F_NORM
        cell.alignment = A_L if kind in ("text", "bold") else A_C
        if right:
            cell.border = Border(left=THIN if left else None, right=THIN, bottom=THIN if bottom else None)
        else:
            cell.border = B_LB if (left and bottom) else B_L if left else B_B if bottom else Border()

    def table_sheet(name, heading, groups, data, freeze="D6"):
        """groups: [(group label or None, [(sub label, width, kind), ...])]. A group with one column and no sub label
        is merged over rows 4-5, like Reference / Name in the workbook. data: list of row lists (one value per column)."""
        ws = book.create_sheet(name)
        cols = [(g, s, w, k, i == 0) for g, subs in groups for i, (s, w, k) in enumerate(subs)]
        title(ws, heading, len(cols) + 1)
        c = 2
        for g, subs in groups:
            n = len(subs)
            for j, (s, w, k) in enumerate(subs):
                ws.column_dimensions[L(c + j)].width = w
            if n == 1 and subs[0][0] is None:
                ws.merge_cells(start_row=4, start_column=c, end_row=5, end_column=c)
                ws.cell(4, c, g).font = F_BOLD; ws.cell(4, c).alignment = A_C
                ws.cell(4, c).border = B_BOX; ws.cell(5, c).border = B_BOX
            else:
                for j, (s, w, k) in enumerate(subs):
                    h = ws.cell(4, c + j, g if j == 0 else None)
                    h.font = F_BOLD; h.alignment = A_CC
                    h.border = Border(left=THIN if j == 0 else None, right=THIN if j == n - 1 else None, top=THIN, bottom=THIN)
                    sh = ws.cell(5, c + j, s); sh.alignment = A_C; sh.border = B_BOX
            c += n
        for i, row in enumerate(data):
            last = i == len(data) - 1
            for j, ((g, s, w, k, first), v) in enumerate(zip(cols, row)):
                put(ws, 6 + i, 2 + j, v, k, left=first, bottom=last, right=(j == len(cols) - 1))
        ws.freeze_panes = freeze
        return ws

    def block_sheet(ws, heading, sign, nblocks, vol_label="Volume (m)"):
        """Lay out (or relabel) a sheet like the 260-day template: Reference / Name / Ticker / GICS Sector, then one
        block of Date / Open / Close / Volume per trading day, numbered from the announcement."""
        title(ws, heading, 5 + 4 * nblocks)
        for c, (lab, w) in enumerate([("Reference", 10.7), ("Name", 38.7), ("Ticker", 10.7), ("GICS\nSector", 23.1)], start=2):
            if ws.cell(4, c).value is None:
                ws.merge_cells(start_row=4, start_column=c, end_row=5, end_column=c)
                ws.cell(4, c, lab).font = F_BOLD; ws.cell(4, c).alignment = A_C
                ws.cell(4, c).border = B_BOX; ws.cell(5, c).border = B_BOX
            ws.column_dimensions[L(c)].width = w
        for k in range(nblocks):
            c0 = 6 + 4 * k
            h = ws.cell(4, c0)
            if h.value is None:
                h.value = k + (1 if sign == "-" else 0)
            h.number_format = f'"Announcement {sign}"0'; h.font = F_BOLD; h.alignment = A_CC
            h.border = Border(left=THIN, top=THIN, bottom=THIN)
            ws.cell(4, c0 + 1).border = ws.cell(4, c0 + 2).border = Border(top=THIN, bottom=THIN)
            ws.cell(4, c0 + 1).alignment = ws.cell(4, c0 + 2).alignment = ws.cell(4, c0 + 3).alignment = A_CC
            ws.cell(4, c0 + 3).border = Border(right=THIN, top=THIN, bottom=THIN)
            for j, (lab, w) in enumerate([("Date", 13.3), ("Open Price", 11.7), ("Close Price", 11.7), (vol_label, 11.7)]):
                sh = ws.cell(5, c0 + j); sh.value = lab; sh.alignment = A_C; sh.border = B_BOX
                ws.column_dimensions[L(c0 + j)].width = w
        ws.freeze_panes = "F6"

    def fill_blocks(ws, row, values, last):
        """values: list of (date, open, close, volume_m) per block, None for a missing day."""
        for k, v in enumerate(values):
            c0 = 6 + 4 * k
            v = v or (None, None, None, None)
            put(ws, row, c0, v[0], "date", left=True, bottom=last)
            for j in range(1, 4):
                put(ws, row, c0 + j, v[j], "num", bottom=last)

    def ident(ws, row, e, last):
        for c, (v, kind) in enumerate([(None, "text"), (e["name"], "bold"), (e["ticker"], "text"), (e["sector"], "text")], start=2):
            put(ws, row, c, v, kind, left=True, bottom=last)
            if c == 4: ws.cell(row, c).alignment = A_C

    # per-event day lookups
    by_ev = {eid: d.drop_duplicates("rel_day").set_index("rel_day") for eid, d in daily.groupby("event_id")}
    def stock_day(eid, rel):
        d = by_ev.get(eid)
        if d is None or rel not in d.index: return None
        x = d.loc[rel]
        return (x["date"], x["open"], x["close"], x["vol"] / 1e6 if pd.notna(x["vol"]) else None)
    def spy_day(ann, rel):
        i = cal.searchsorted(ann) + rel
        if i < 0 or i >= len(cal): return None
        day = cal[i]; x = spy.loc[day]
        return (day, x["open"], x["close"], x["vol"] / 1e6 if pd.notna(x["vol"]) else None)

    order = ev.sort_values("row")                     # the workbook's own row order (newest first)
    last_row = order.row.max()

    # 7a. 260-day sheets (days -1 to -260)
    for sn, src_fn, heading in [(HIST, "stock", "Additions 260-day Historical Prices"), (HIST_SPY, "spy", "SPY 260-day Historical Prices")]:
        ws = book[sn] if sn in book.sheetnames else book.create_sheet(sn)
        block_sheet(ws, heading, "-", HIST_DAYS)
        for _, e in order.iterrows():
            r = int(e.row)
            if ws.cell(r, 3).value not in (None, e["name"]):
                print(f"  {sn}: row {r} is '{ws.cell(r, 3).value}', expected '{e['name']}'; written by the Clean sheet's row")
            ident(ws, r, e, r == last_row)
            vals = [stock_day(e.event_id, -k) if src_fn == "stock" else spy_day(e.ann, -k) for k in range(1, HIST_DAYS + 1)]
            fill_blocks(ws, r, vals, r == last_row)
        print(f"Filled {sn}")

    # 7b. event-window sheets (day 0 = announcement, to the inclusion day: the strategy's holding period)
    span_days = {e.event_id: int(cal.searchsorted(e.inc) - cal.searchsorted(e.ann)) for _, e in ev.iterrows()}
    nwin = max(span_days.values()) + 1
    for sn, src_fn, heading in [(WIN, "stock", "Additions Event Window: Announcement to Inclusion"),
                                (WIN_SPY, "spy", "SPY Event Window: Announcement to Inclusion")]:
        ws = book.create_sheet(sn)
        block_sheet(ws, heading, "+", nwin)
        for i, (_, e) in enumerate(order.iterrows()):
            r = 6 + i
            ident(ws, r, e, i == len(order) - 1)
            n = span_days[e.event_id] + 1
            vals = [(stock_day(e.event_id, k) if src_fn == "stock" else spy_day(e.ann, k)) if k < n else None for k in range(nwin)]
            fill_blocks(ws, r, vals, i == len(order) - 1)
        print(f"Filled {sn}")

    # 7c. replace prices in the Clean and All sheets (red and blue rows untouched)
    PRICE_CELLS = [(7, "ann_c"), (8, "ann_v"), (10, "a1_o"), (11, "a1_v"), (13, "inc_o"), (14, "inc_c"), (15, "inc_v"),
                   (17, "i1_c"), (18, "i1_v"), (20, "i3_c"), (21, "i3_v")]
    def is_black(cell):
        c = cell.font.color
        if c is None: return True
        if c.type == "theme": return c.theme in (1, None)
        rgb = (c.rgb or "")[-6:].upper()
        return rgb in ("", "000000")
    def set_prices(ws, r, e, prefix):
        """Overwrite the price and volume cells; a cell we have no value for keeps its old value (counted)."""
        kept = 0
        for col, key in PRICE_CELLS:
            v = e[prefix + key]
            cell = ws.cell(r, col)
            if pd.isna(v):
                kept += cell.value is not None
                continue
            cell.value = float(v / 1e6 if key.endswith("_v") else v)
            cell.number_format = "0.00"
        return kept
    def set_note(ws, r, text):
        old = ws.cell(r, 25).value
        parts = [p for p in str(old).split(" | ") if p and not p.startswith("Prices: ")] if old else []
        ws.cell(r, 25).value = " | ".join(parts + [text])
    def note_for(e, kept):
        s = f"Prices: CRSP (PERMNO {int(e.permno)})" if e.price_source == "CRSP" and pd.notna(e.permno) else \
            "Prices: Yahoo (after CRSP's last date)" if e.price_source == "Yahoo" else "Prices: not found in CRSP"
        return s + (f", {kept} cells kept from before (no new value)" if kept else "")

    for sn in [CLEAN_SPY, ALL_SPY]:
        for c in (8, 11, 15, 18, 21):
            if book[sn].cell(5, c).value == "Volume": book[sn].cell(5, c).value = "Volume (m)"
    for _, e in ev.iterrows():
        r = int(e.row)
        kept = set_prices(book[CLEAN], r, e, ""); set_prices(book[CLEAN_SPY], r, e, "spy_"); set_note(book[CLEAN], r, note_for(e, kept))
    all_ws = book[ALL]
    all_rows = {}
    for r in range(6, all_ws.max_row + 1):
        nm, d = all_ws.cell(r, 3).value, all_ws.cell(r, 6).value
        if all_ws.cell(r, 2).value is None and nm and isinstance(d, (dt.date, dt.datetime)) and is_black(all_ws.cell(r, 3)):
            all_rows.setdefault((str(nm).strip(), pd.Timestamp(d).normalize()), r)
    missing_all = []
    for _, e in ev.iterrows():
        r = all_rows.get((str(e["name"]).strip(), e.ann.normalize()))
        if r is None:
            missing_all.append(e["name"]); continue
        kept = set_prices(all_ws, r, e, ""); set_prices(book[ALL_SPY], r, e, "spy_"); set_note(all_ws, r, note_for(e, kept))
    if missing_all:
        print("Not found in the All sheet (left as they were):", ", ".join(missing_all))

    # 7d. Event Data
    ed = [[e.event_id, e["name"], e.ticker, e.sector,
           e.price_source, e.permno, e.src_name, e.match_method,
           e.ann, e.a1, e.inc, e.gap_days, "Yes" if e.quarterly else "No",
           e.mktcap_ann_musd, e.adv_dollar_musd, e.demand_days, e.n_est_days,
           e.pre_inc, e.pre_inc_c, e.spy_pre_inc_c,
           e.div_yield_window, "Yes" if e.split_in_window else "No", e.chk_A_diff, e.chk_B_diff,
           "; ".join(x for x in [e.review, e.quarterly_note] if x)] for _, e in order.iterrows()]
    table_sheet(EVD, "Event Data", [
        ("Reference", [(None, 10.7, "text")]), ("Name", [(None, 38.7, "bold")]), ("Ticker", [(None, 10.7, "text")]),
        ("GICS\nSector", [(None, 23.1, "text")]),
        ("Price Data Match", [("Source", 9.0, "text"), ("PERMNO", 10.0, "int"), ("Name at source", 34.0, "text"), ("Matched by", 30.0, "text")]),
        ("Dates", [("Announcement", 13.3, "date"), ("Announc. +1", 13.3, "date"), ("Inclusion", 13.3, "date"),
                   ("Gap (cal. days)", 13.0, "int"), ("Quarterly rebal.", 13.0, "text")]),
        ("Size & Liquidity", [("Mkt cap ($m)", 13.0, "num"), ("Avg $ vol. ($m)", 13.0, "num"),
                              ("Demand days", 12.0, "num"), ("Beta days", 10.0, "int")]),
        ("Day Before Inclusion", [("Date", 13.3, "date"), ("Close", 11.0, "num"), ("SPY close", 11.0, "num")]),
        ("Checks", [("Dividend yield in window", 13.0, "pct"), ("Split in window", 10.0, "text"),
                    ("vs old: window A", 12.0, "pct"), ("vs old: window B", 12.0, "pct")]),
        ("Notes", [(None, 60.0, "text")]),
    ], ed)

    # 7e. Fama-French
    table_sheet(FFS, "Fama-French Daily Factors (US)", [
        ("Date", [(None, 13.3, "date")]),
        ("Factors (daily, %)", [("Mkt-RF", 11.0, "pct3"), ("SMB", 11.0, "pct3"), ("HML", 11.0, "pct3"), ("RF", 11.0, "pct3"), ("Mom", 11.0, "pct3")]),
    ], [[d] + [ff.loc[d].get(c, np.nan) for c in ["Mkt-RF", "SMB", "HML", "RF", "Mom"]] for d in ff.index], freeze="C6")

    # 7f. Needs Review
    rv = order[order.review != ""]
    table_sheet(REV, "Needs Review", [
        ("Reference", [(None, 10.7, "text")]), ("Name", [(None, 38.7, "bold")]), ("Ticker", [(None, 10.7, "text")]),
        ("Dates", [("Announcement", 13.3, "date"), ("Inclusion", 13.3, "date")]),
        ("Price Data Match", [("Source", 9.0, "text"), ("PERMNO", 10.0, "int"), ("Name at source", 34.0, "text"), ("Matched by", 30.0, "text")]),
        ("Issue", [(None, 80.0, "text")]),
    ], [[e.event_id, e["name"], e.ticker, e.ann, e.inc, e.price_source, e.permno, e.src_name, e.match_method, e.review]
        for _, e in rv.iterrows()] or [["", "Nothing to review", "", None, None, "", None, "", "", ""]])

    # 7g. Notes & Sources
    srcs = "; ".join(f"crsp.{s['table']} to {s['last'].date()}" for s in sources)
    notes = [
        ("Built", f"{dt.datetime.now():%Y-%m-%d %H:%M} from {Path(SRC).name} (additions read from '{CLEAN}')"),
        ("Additions", f"{len(ev)}: {int((ev.price_source == 'CRSP').sum())} priced from CRSP, {int((ev.price_source == 'Yahoo').sum())} from Yahoo"),
        ("Stock prices", f"CRSP daily stock file via WRDS ({srcs}). Additions whose window ends after {crsp_last.date()} come from Yahoo Finance in full."),
        ("Benchmark", f"SPY (CRSP PERMNO {SPY_PERMNO}) to {crsp_last.date()}, then Yahoo Finance to {spy.index.max().date()}."),
        ("Factors", f"Kenneth French data library, US daily 3 factors and momentum, {ff.index.min().date()} to {ff.index.max().date()}."),
        ("Prices", "USD, not adjusted for dividends. Adjusted for splits to the announcement day's share basis, so they equal "
                   "CRSP's own prices on and after the announcement day."),
        ("Volume", "Millions of shares, on the same share basis as the prices."),
        ("Matching", "CRSP permanent ID (PERMNO) from the ticker valid on the announcement date, then the inclusion date, "
                     "then a company-name search; the CRSP name is compared with ours (Event Data, 'Name at source')."),
        ("260-day sheets", "Block k = k trading days before the announcement (Announcement -1 is the day before). "
                           "Trading days are the days SPY traded. Beta uses days -250 to -30; average volume uses -60 to -11."),
        ("Event Window sheets", "Block 0 = announcement day, through the inclusion day (the last block on each row)."),
        ("Clean / All sheets", "Price and volume cells replaced with the values above; the Notes column says the source. "
                               "Red and blue rows are unchanged. SPY sheets now hold SPY, not the S&P 500 index."),
        ("Event Data: Demand days", "Market cap / average daily dollar volume (days -60 to -11): days of normal trading "
                                    "needed to absorb the index funds' purchase."),
        ("Event Data: Quarterly rebal.", "Part of S&P's scheduled quarterly rebalance: effective on the Monday after the third "
                                         "Friday of Mar/Jun/Sep/Dec (Tuesday if that Monday is a holiday), announced at least "
                                         "7 days before, and not a replacement for a company leaving because of a takeover, "
                                         "spinoff or conversion. Checked against S&P's press releases (7 corrected, see Notes)."),
        ("Event Data: Dividend yield", "Dividends with an ex-date after Ann+1's open and up to the inclusion day, as % of price."),
        ("Event Data: vs old", "Our window return minus the one from the prices previously in the workbook (should be about 0)."),
        ("Fama-French sheet", "Daily returns in % (0.01% = 0.0001). RF is the one-month T-bill rate per day."),
    ]
    table_sheet(NOTES, "Notes & Sources", [("Item", [(None, 30.0, "bold")]), ("Detail", [(None, 120.0, "text")])],
                [[a, b] for a, b in notes], freeze=None)

    book.save(OUT)
    print("Saved", OUT)
    return Path(OUT)


# =====================================================================================================================
# PART 2. CALCULATIONS: every statistic of the workbook, recomputed from the raw data sheets
# =====================================================================================================================

# ---------------------------------------------------------------------------------------------------------------------
# 2.0 Small helpers that copy Excel's functions exactly
# ---------------------------------------------------------------------------------------------------------------------
def is_number(v):
    """Excel's ISNUMBER: a typed number (True/False and text do not count)."""
    return isinstance(v, (int, float, np.integer, np.floating)) and not isinstance(v, (bool, np.bool_))


def num(v):
    """A typed number as float, anything else (blank, text, a formula string) as NaN."""
    return float(v) if is_number(v) else np.nan


def day(v):
    """A typed date as numpy datetime64[D], anything else as NaT."""
    if isinstance(v, dt.datetime):
        return np.datetime64(v.date(), "D")
    if isinstance(v, dt.date):
        return np.datetime64(v, "D")
    return np.datetime64("NaT", "D")


def to_date(d):
    """numpy/pandas date -> datetime.date (for printing and comparing)."""
    return pd.Timestamp(d).date()


def t_dist_2t(x, df):
    """T.DIST.2T(x, df). Excel and LibreOffice truncate df to an integer (matters only for the Welch test)."""
    return 2.0 * stats.t.sf(abs(x), math.floor(df))


def t_inv_2t(prob, df):
    """T.INV.2T(prob, df): the t value that leaves prob/2 in each tail."""
    return stats.t.ppf(1.0 - prob / 2.0, math.floor(df))


def norm_2t(z):
    """2 * (1 - NORM.S.DIST(|z|, TRUE)), written exactly as on the sheets (it is exactly 0 for |z| > 8.3)."""
    return 2.0 * (1.0 - stats.norm.cdf(abs(z)))


def rank_avg(x):
    """RANK.AVG(value, range, 1): ascending ranks, ties get the average rank."""
    return stats.rankdata(x, method="average")


def stdev_s(x):
    """STDEV.S: sample standard deviation (divides by n - 1)."""
    return float(np.std(np.asarray(x, float), ddof=1))


def text_pct(x, decimals=0):
    """Excel TEXT(x, "0%") / TEXT(x, "0.0%"), used in a few labels."""
    return f"{x * 100:.{decimals}f}%"


def text_dmy(d):
    """Excel TEXT(date, "d mmm yyyy"), e.g. 29 Dec 2017."""
    d = to_date(d)
    return f"{d.day} {d:%b %Y}"


# ---------------------------------------------------------------------------------------------------------------------
# 2.1 Read the raw data sheets (typed inputs only)
# ---------------------------------------------------------------------------------------------------------------------
def find_workbook(arg=None):
    """The workbook to use: the path given on the command line, else the newest 'SPX Additions v<N>.xlsx' in this
    script's folder, else in ../../data/prices next to it (the project layout). Versions compare as integers."""
    if arg:
        p = Path(arg).expanduser()
        if not p.exists():
            sys.exit(f"Workbook not found: {p}")
        return p
    here = Path(__file__).resolve().parent
    for folder in (here, here.parent.parent / "data" / "prices"):
        found = []
        for p in folder.glob("SPX Additions v*.xlsx"):
            m = re.fullmatch(r"SPX Additions v(\d+)\.xlsx", p.name)
            if m:
                found.append((int(m.group(1)), p))
        if found:
            return max(found)[1]
    sys.exit("No 'SPX Additions v<N>.xlsx' found next to this script or in ../../data/prices. "
             "Give the workbook's path: python spx_additions_analysis.py path/to/workbook.xlsx")


def load_inputs(path):
    """Read every input the calculations need, from the raw data sheets only.

    The workbook is opened with data_only=False, so a typed value comes back as a number, date or text and a formula
    comes back as a string starting with '='. Every input below is a typed value. The only formulas inside the data
    sheets that matter are Fama-French Daily F4701:F4714 ('=$F$4700': RF for 1-21 Sep 2026, not yet published by
    French, repeats 31 Aug's value); that rule is reproduced in daily_account(), never read.
    """
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=False)

    def rows(sheet, r0, r1, max_col=None):
        return [list(r) for r in wb[sheet].iter_rows(min_row=r0, max_row=r1, min_col=1, max_col=max_col,
                                                     values_only=True)]

    # Event rows: from row 6 down to the last row with a name (C) and an announcement date (F) in the Clean sheet.
    # The rows below them hold a colour key. Every per-event sheet has the same event in the same row.
    clean_all = rows(CLEAN, FIRST_ROW, None, 14)
    n = 0
    while n < len(clean_all) and clean_all[n][2] is not None and isinstance(clean_all[n][5], dt.datetime):
        n += 1
    last_row = FIRST_ROW + n - 1
    clean = clean_all[:n]
    cspy = rows(CLEAN_SPY, FIRST_ROW, last_row, 14)
    evd = rows(EVENT_DATA, FIRST_ROW, last_row, 21)
    aew = rows(WIN, FIRST_ROW, last_row)
    sew = rows(WIN_SPY, FIRST_ROW, last_row)
    a260 = rows(HIST, FIRST_ROW, last_row, 5 + 4 * HIST_DAYS)
    s260 = rows(HIST_SPY, FIRST_ROW, last_row, 5 + 4 * HIST_DAYS)
    ffr = rows(FFS, FIRST_ROW, None, 7)
    wb.close()

    # the same company must sit in the same row of every per-event sheet (column C = name)
    for sheet, data in [(CLEAN_SPY, cspy), (EVENT_DATA, evd), (WIN, aew), (WIN_SPY, sew), (HIST, a260),
                        (HIST_SPY, s260)]:
        bad = [FIRST_ROW + i for i in range(n) if data[i][2] != clean[i][2]]
        if bad:
            sys.exit(f"Row mismatch: '{sheet}' rows {bad[:5]} do not hold the same company as '{CLEAN}'")

    col = lambda data, letter: [r[ord(letter) - 65] for r in data]          # single-letter columns A..Z
    ev = pd.DataFrame({
        "row": np.arange(FIRST_ROW, last_row + 1),
        "ref": col(evd, "B"), "name": col(clean, "C"), "ticker": col(clean, "D"),
        "announced": pd.to_datetime([day(v) for v in col(clean, "F")]),
        "included": pd.to_datetime([day(v) for v in col(clean, "L")]),
        "quarterly": col(evd, "N"),
        # stock prices: announcement close (G), Ann+1 open = buy (J), inclusion open = sell A (M), close = sell B (N),
        # close the day before inclusion (Event Data T)
        "s_ann_close": [num(v) for v in col(clean, "G")], "s_buy_open": [num(v) for v in col(clean, "J")],
        "s_pre_close": [num(v) for v in col(evd, "T")],
        "s_inc_open": [num(v) for v in col(clean, "M")], "s_inc_close": [num(v) for v in col(clean, "N")],
        # SPY on the same dates (Clean SPY sheet; Event Data U for the day before inclusion)
        "m_ann_close": [num(v) for v in col(cspy, "G")], "m_buy_open": [num(v) for v in col(cspy, "J")],
        "m_pre_close": [num(v) for v in col(evd, "U")],
        "m_inc_open": [num(v) for v in col(cspy, "M")], "m_inc_close": [num(v) for v in col(cspy, "N")],
    })

    def blocks(data, nk, field):
        """The j-th block of a row (j = 0, 1, ...) starts in the 0-based column 5 + 4j with date, open, close, volume.
        field: 0 = date, 1 = open, 2 = close. Returns an (events x nk) array."""
        out = np.full((n, nk), np.nan) if field else np.full((n, nk), np.datetime64("NaT", "D"))
        for i, r in enumerate(data):
            for j in range(nk):
                c = 5 + 4 * j + field
                v = r[c] if c < len(r) else None
                out[i, j] = num(v) if field else day(v)
        return out

    # Event Window sheets: block k = k trading days after the announcement (k = 0 is the announcement day), in the
    # 1-based columns 6+4k .. 9+4k (block 0 = F:I, block 1 = J:M). Blocks after the inclusion day are empty.
    nwin = (max(len(r) for r in aew) - 5) // 4
    win = SimpleNamespace(date=blocks(aew, nwin, 0), open=blocks(aew, nwin, 1), close=blocks(aew, nwin, 2),
                          spy_date=blocks(sew, nwin, 0), spy_open=blocks(sew, nwin, 1), spy_close=blocks(sew, nwin, 2))
    # 260-day sheets: block k = k trading days BEFORE the announcement (k = 1..260), in the 1-based columns
    # 4k+2 .. 4k+5 (block 1 = F:I). Column 0 of these arrays is left empty so that column k holds block k.
    pad_num, pad_day = np.full((n, 1), np.nan), np.full((n, 1), np.datetime64("NaT", "D"))
    hist = SimpleNamespace(close=np.hstack([pad_num, blocks(a260, HIST_DAYS, 2)]),
                           spy_date=np.hstack([pad_day, blocks(s260, HIST_DAYS, 0)]),
                           spy_close=np.hstack([pad_num, blocks(s260, HIST_DAYS, 2)]))

    # Fama-French Daily: B date, C Mkt-RF, D SMB, E HML, F RF, G Mom (decimals). Typed rows only.
    ff = pd.DataFrame([{"date": day(r[1]), "mkt": num(r[2]), "smb": num(r[3]), "hml": num(r[4]), "rf": num(r[5]),
                        "mom": num(r[6])} for r in ffr if isinstance(r[1], dt.datetime)])
    return SimpleNamespace(path=Path(path), n=n, ev=ev, win=win, hist=hist, ff=ff)


def spy_close_by_date(raw):
    """SPY's close for every date found in the SPY 260-day and SPY Event Window sheets (the Daily Account links each
    day to one of these cells). Checks that no date has two different closes."""
    dates = np.concatenate([raw.hist.spy_date[:, 1:].ravel(), raw.win.spy_date.ravel()])
    closes = np.concatenate([raw.hist.spy_close[:, 1:].ravel(), raw.win.spy_close.ravel()])
    ok = ~np.isnat(dates) & ~np.isnan(closes)
    s = pd.Series(closes[ok], index=pd.to_datetime(dates[ok]))
    conflicts = s.groupby(level=0).nunique()
    if (conflicts > 1).any():
        sys.exit(f"SPY sheets disagree on the close for {list(conflicts[conflicts > 1].index[:5])}")
    return s.groupby(level=0).first().sort_index()


# ---------------------------------------------------------------------------------------------------------------------
# 2.2 Trade Returns and Beta Estimation sheets
# ---------------------------------------------------------------------------------------------------------------------
def beta_estimation(raw):
    """Beta Estimation sheet: for each stock, the OLS slope of its daily returns on SPY's over days -250 to -30
    before the announcement. Return day -k = close(day -k) / close(day -k-1) - 1 when both closes exist. A stock
    with fewer than MIN_BETA_DAYS returns gets beta = 1. SLOPE only uses days where both returns exist."""
    k = np.arange(BETA_FIRST_DAY, BETA_LAST_DAY - 1, -1)                  # 250, 249, ..., 30
    stock = raw.hist.close[:, k] / raw.hist.close[:, k + 1] - 1.0
    spy = raw.hist.spy_close[:, k] / raw.hist.spy_close[:, k + 1] - 1.0
    days = (~np.isnan(stock)).sum(axis=1)                                  # COUNT of stock returns (column D)
    slope = np.full(raw.n, np.nan)
    for i in range(raw.n):
        both = ~np.isnan(stock[i]) & ~np.isnan(spy[i])
        x, y = spy[i, both], stock[i, both]
        if len(x) > 1:
            slope[i] = ((x - x.mean()) * (y - y.mean())).sum() / ((x - x.mean()) ** 2).sum()
    used = np.where(days >= MIN_BETA_DAYS, slope, 1.0)                     # column G
    return pd.DataFrame({"days": days, "slope": np.where(days >= MIN_BETA_DAYS, slope, np.nan), "beta": used})


def trade_returns(raw, beta):
    """Trade Returns sheet: one row per trade (workbook column letters in the comments)."""
    ev = raw.ev
    tr = ev[["row", "ref", "name", "ticker", "announced", "included", "quarterly"]].copy()
    tr["year"] = tr.announced.dt.year                                                        # G
    s_ann1_close, m_ann1_close = raw.win.close[:, 1], raw.win.spy_close[:, 1]                # K, Q (Ann+1 close)
    tr["R_A"] = ev.s_inc_open / ev.s_buy_open - 1                                            # U stock, Window A
    tr["M_A"] = ev.m_inc_open / ev.m_buy_open - 1                                            # V SPY, Window A
    tr["AR_A"] = tr.R_A - tr.M_A                                                             # W
    tr["R_B"] = ev.s_inc_close / ev.s_buy_open - 1                                           # X stock, Window B
    tr["M_B"] = ev.m_inc_close / ev.m_buy_open - 1                                           # Y SPY, Window B
    tr["AR_B"] = tr.R_B - tr.M_B                                                             # Z
    tr["beta"] = beta.beta.values                                                            # AA
    tr["AR_A_adj"] = tr.R_A - tr.beta * tr.M_A                                               # AB risk-adjusted A
    tr["AR_B_adj"] = tr.R_B - tr.beta * tr.M_B                                               # AC risk-adjusted B
    tr["V5_stock"] = ev.s_pre_close / ev.s_buy_open - 1                                      # AD V5: sell the day
    tr["V5_spy"] = ev.m_pre_close / ev.m_buy_open - 1                                        # AE before inclusion
    tr["V5"] = tr.V5_stock - tr.V5_spy                                                       # AF
    # Figure 1: abnormal return since the announcement close, at each moment (AG..AL)
    tr["CAR_ann_close"] = ev.s_ann_close / ev.s_ann_close - ev.m_ann_close / ev.m_ann_close  # AG (= 0)
    tr["CAR_buy_open"] = ev.s_buy_open / ev.s_ann_close - ev.m_buy_open / ev.m_ann_close     # AH overnight jump
    tr["CAR_ann1_close"] = s_ann1_close / ev.s_ann_close - m_ann1_close / ev.m_ann_close     # AI
    tr["CAR_pre_close"] = ev.s_pre_close / ev.s_ann_close - ev.m_pre_close / ev.m_ann_close  # AJ
    tr["CAR_inc_open"] = ev.s_inc_open / ev.s_ann_close - ev.m_inc_open / ev.m_ann_close     # AK
    tr["CAR_inc_close"] = ev.s_inc_close / ev.s_ann_close - ev.m_inc_close / ev.m_ann_close  # AL
    # AM holding days: the Event Window block whose date is the inclusion date
    hit = raw.win.date == ev.included.values.astype("datetime64[D]")[:, None]
    if not hit.any(axis=1).all():
        sys.exit("Some inclusion dates are not in the Event Window sheet")
    tr["days"] = hit.argmax(axis=1)
    tr["label"] = [f"{nm} ({y})" for nm, y in zip(tr.name, tr.year)]                        # AR
    # The prices (I..T), also used by the worked examples of Tests 1 and 2
    for c in ["s_ann_close", "s_buy_open", "s_pre_close", "s_inc_open", "s_inc_close",
              "m_ann_close", "m_buy_open", "m_pre_close", "m_inc_open", "m_inc_close"]:
        tr[c] = ev[c].values
    tr["s_ann1_close"], tr["m_ann1_close"] = s_ann1_close, m_ann1_close
    return tr


# ---------------------------------------------------------------------------------------------------------------------
# 2.3 Holding Days and Daily Account sheets
# ---------------------------------------------------------------------------------------------------------------------
def holding_days(raw, tr):
    """Holding Days sheet: one row per position per trading day held (day 1 = Ann+1, day E = inclusion day),
    positions in sheet order bottom-up (oldest first), with each day's piece of the return.

    Window A stock: day 1 open->close, day E previous close->open (sold at the open), other days close->close.
    Window B stock: day 1 open->close, later days close->close (sold at the close). SPY pieces the same way.
    Long + SPY account, Window A: day 1 = SPY previous close->open, then stock open->close (the account sits in SPY
    until it buys); day E = stock previous close->open, then SPY open->close; other days stock close->close.
    Long + SPY account, Window B: day 1 as in A, later days stock close->close.
    Also adds to tr the chain checks of Trade Returns AN..AQ.
    """
    w = raw.win
    parts = []
    for i in reversed(range(raw.n)):
        E = int(tr.days.iat[i])
        d = np.arange(1, E + 1)
        so, sc, sp = w.open[i, 1:E + 1], w.close[i, 1:E + 1], w.close[i, 0:E]           # stock open, close, prev close
        mo, mc, mp = w.spy_open[i, 1:E + 1], w.spy_close[i, 1:E + 1], w.spy_close[i, 0:E]
        first, last = d == 1, (d == E) & (d != 1)                                         # day 1 is tested first
        stock_A = np.where(first, sc / so - 1, np.where(last, so / sp - 1, sc / sp - 1))
        spy_A = np.where(first, mc / mo - 1, np.where(last, mo / mp - 1, mc / mp - 1))
        stock_B = np.where(first, sc / so - 1, sc / sp - 1)
        spy_B = np.where(first, mc / mo - 1, mc / mp - 1)
        long_A = np.where(first, (mo / mp) * (sc / so) - 1, np.where(last, (so / sp) * (mc / mo) - 1, sc / sp - 1))
        long_B = np.where(first, (mo / mp) * (sc / so) - 1, sc / sp - 1)
        parts.append(pd.DataFrame({
            "trade": i, "row": tr.row.iat[i], "name": tr.name.iat[i], "day": d, "date": w.date[i, 1:E + 1],
            "stock_open": so, "stock_close": sc, "stock_prev_close": sp,
            "spy_open": mo, "spy_close": mc, "spy_prev_close": mp,
            "stock_A": stock_A, "spy_A": spy_A, "ar_A": stock_A - spy_A,
            "stock_B": stock_B, "spy_B": spy_B, "ar_B": stock_B - spy_B,
            "long_spy_A": long_A, "long_spy_B": long_B}))
    hd = pd.concat(parts, ignore_index=True)
    hd["date"] = pd.to_datetime(hd.date)
    # Trade Returns AN..AQ: chaining a trade's daily pieces must give back its window return
    chain = hd.groupby("trade").agg(
        sA=("stock_A", lambda s: np.prod(1 + s.values) - 1), sB=("stock_B", lambda s: np.prod(1 + s.values) - 1),
        mA=("spy_A", lambda s: np.prod(1 + s.values) - 1), mB=("spy_B", lambda s: np.prod(1 + s.values) - 1))
    tr["gap_stock_A"] = chain.sA.values - tr.R_A.values                                       # AN
    tr["gap_stock_B"] = chain.sB.values - tr.R_B.values                                       # AO
    tr["gap_spy_A"] = chain.mA.values - tr.M_A.values                                         # AP
    tr["gap_spy_B"] = chain.mB.values - tr.M_B.values                                         # AQ
    return hd


def daily_account(raw, hd, spy):
    """Daily Account sheet: one row per trading day from the first Fama-French day of 2010 to the last inclusion.

    B1 = SPY bought and held. Long + SPY (Window A or B) = on days with open positions, the equal-weighted average
    of the positions' pieces that day (Holding Days); on days with none, SPY (so it equals B1 that day).
    Excess return = return - RF. Value of $1 = cumulative product; drawdown = value / running maximum - 1.
    """
    ff = raw.ff
    has_f = ~np.isnan(ff.mkt.values)
    ff_last = ff.date.values[int(has_f.sum()) - 1]          # Settings C22: INDEX(dates, COUNT(Mkt-RF))
    last_inclusion = raw.ev.included.max().to_datetime64().astype("datetime64[D]")
    dates = ff.date.values[ff.date.values >= np.datetime64(ACCOUNT_START)]
    if dates.max() < last_inclusion:
        # A workbook straight from the data collection has no rows after French's last day: the account then
        # continues on SPY's trading days up to the last inclusion day (the builder added exactly these days).
        extra = spy.index.values.astype("datetime64[D]")
        dates = np.r_[dates, extra[(extra > dates.max()) & (extra <= last_inclusion)]]
    rf_by_date = pd.Series(ff.rf.values, index=pd.to_datetime(ff.date.values))
    rf_last = rf_by_date[pd.Timestamp(ff_last)]
    da = pd.DataFrame({"date": pd.to_datetime(dates)})
    da["rf"] = [rf_by_date.get(d, np.nan) if d <= pd.Timestamp(ff_last) else rf_last for d in da.date]  # C
    held = hd.groupby("date")
    da["positions"] = held.size().reindex(da.date).fillna(0).astype(int).values                       # D
    da["spy_close"] = spy.reindex(da.date).values                                                      # E
    da["price_found"] = ~np.isnan(da.spy_close.values)                                                 # H
    first_prev = spy[spy.index < da.date.iat[0]].iloc[-1]                                              # 31 Dec 2009
    da["spy_prev_close"] = np.r_[first_prev, da.spy_close.values[:-1]]                                 # F
    da["B1"] = da.spy_close / da.spy_prev_close - 1                                                    # G
    da["avg_A"] = held.long_spy_A.mean().reindex(da.date).values                                       # I
    da["avg_B"] = held.long_spy_B.mean().reindex(da.date).values                                       # K
    on = da.positions.values > 0
    da["A"] = np.where(on, da.avg_A.values, da.B1.values)                                              # J
    da["B"] = np.where(on, da.avg_B.values, da.B1.values)                                              # L
    for k in ["A", "B1", "B"]:
        da["x_" + k] = da[k] - da.rf                                                                   # M, N, O
        da["value_" + k] = np.cumprod(1 + da[k].values)                                                # P, Q, R
        v = da["value_" + k].values
        da["dd_" + k] = v / np.maximum.accumulate(v) - 1                                               # S, T, U
    # Fama-French factors on each day (blank after French's last day)
    f = ff.set_index(pd.to_datetime(ff.date.values))
    for c in ["mkt", "smb", "hml", "mom"]:
        da[c] = f[c].reindex(da.date).values
    da.attrs["ff_last"] = pd.Timestamp(ff_last)
    return da


def day_number(dates, d):
    """MATCH(d, dates, 1): 1-based position of the last date on or before d."""
    return int(np.searchsorted(dates.values.astype("datetime64[D]"), np.datetime64(d, "D"), side="right"))


# ---------------------------------------------------------------------------------------------------------------------
# 2.4 Regression Workings: OLS with Newey-West standard errors
# ---------------------------------------------------------------------------------------------------------------------
def newey_west_ols(y, X, lags=NW_LAGS):
    """b = (X'X)^-1 X'y; Newey-West covariance with Bartlett weights 1 - l/(lags+1), no small-sample scaling:
    V = (X'X)^-1 [G0 + sum_l w_l (G_l + G_l')] (X'X)^-1 with G_l = sum_t u_t u_{t-l}', u_t = e_t x_t."""
    y, X = np.asarray(y, float), np.asarray(X, float)
    xtx_inv = np.linalg.inv(X.T @ X)
    b = xtx_inv @ (X.T @ y)
    u = X * (y - X @ b)[:, None]
    S = u.T @ u
    for lag in range(1, lags + 1):
        G = u[lag:].T @ u[:-lag]
        S += max(0.0, 1 - lag / (lags + 1)) * (G + G.T)
    se = np.sqrt(np.diag(xtx_inv @ S @ xtx_inv))
    return b, se


def regression(da, ycol, first_day, last_day, four_factor=False):
    """Regress the daily excess return ycol on Mkt-RF (CAPM) or Mkt-RF, SMB, HML, Mom (four-factor) over day
    numbers first_day..last_day (1-based, inclusive). Returns the coefficients, Newey-West SEs and T."""
    s = da.iloc[first_day - 1:last_day]
    cols = ["mkt", "smb", "hml", "mom"] if four_factor else ["mkt"]
    X = np.column_stack([np.ones(len(s))] + [s[c].values for c in cols])
    b, se = newey_west_ols(s[ycol].values, X)
    return SimpleNamespace(b=b, se=se, T=len(s), alpha_year=b[0] * DAYS_PER_YEAR, t=b[0] / se[0])


# ---------------------------------------------------------------------------------------------------------------------
# 2.5 The statistics used on the Test sheets (as laid out in Test 1-2 Workings and Test 4 Workings)
# ---------------------------------------------------------------------------------------------------------------------
def stat_block(x, dates):
    """Every statistic of Tests 1, 2 and 4 for one series of trade returns x with announcement dates.

    Clustered by date: each announcement date counts once (the average of its trades), G = number of dates.
    Wilcoxon signed rank: exact zeros dropped, average ranks of |x|, normal approximation without tie or
    continuity correction. Sign test: exact binomial, two-sided.
    """
    x = np.asarray(x, float)
    dates = np.asarray(dates)
    N = len(x)
    mean, s = x.mean(), stdev_s(x)
    se = s / math.sqrt(N)
    t = mean / se
    groups = pd.Series(x).groupby(dates)
    by_date = groups.mean()                                    # one value per announcement date
    G = len(by_date)
    meanG, sdG = by_date.mean(), stdev_s(by_date.values)
    seG = sdG / math.sqrt(G)
    tG = meanG / seG
    t_star = t_inv_2t(ALPHA, G - 1)
    nz = x[x != 0]
    n = len(nz)
    ranks = rank_avg(np.abs(nz))
    W = ranks[nz > 0].sum()
    EW, VW = n * (n + 1) / 4, n * (n + 1) * (2 * n + 1) / 24
    z = (W - EW) / math.sqrt(VW)
    k = int((x > 0).sum())
    return SimpleNamespace(
        N=N, mean=mean, s=s, se=se, t=t, p=t_dist_2t(t, N - 1),
        G=G, multi=int((groups.size() > 1).sum()), meanG=meanG, sdG=sdG, seG=seG, tG=tG, pG=t_dist_2t(tG, G - 1),
        t_star=t_star, lo=meanG - t_star * seG, hi=meanG + t_star * seG,
        median=float(np.median(x)), n=n, W=W, EW=EW, VW=VW, z=z, pW=norm_2t(z),
        k=k, beat=k / n, pS=min(1.0, 2 * stats.binom.cdf(min(k, n - k), n, 0.5)),
        mde=MDE_MULT * seG)


def sig_word(p, m):
    """'Significant at the 5% level?' cells."""
    return ("Yes, above zero" if m > 0 else "Yes, below zero") if p < ALPHA else "No"


def verdict(b):
    """Test 1's rule (also used in Tests 2 and 4 and Extra Checks E3)."""
    if b.meanG > 0 and b.pG < ALPHA and b.lo > 0 and b.median > 0 and b.pW < ALPHA and b.z > 0:
        return "Works"
    if b.meanG > 0 and b.pG < ALPHA:
        return "Mixed"
    if (b.meanG <= 0 and b.median <= 0) or (b.meanG < 0 and b.pG < ALPHA) or (b.pG >= ALPHA and b.mde <= MIN_EFFECT):
        return "Doesn't work"
    return "Inconclusive"


def combine(a, b):
    """Tests 1 and 2: one answer from the two windows."""
    if "Works" in (a, b):
        return "Works"
    if a == b == "Doesn't work":
        return "Doesn't work"
    if "Mixed" in (a, b):
        return "Mixed"
    return "Inconclusive"


def yes(cond):
    return "Yes" if cond else "No"


def sharpe_daily(x):
    x = np.asarray(x, float)
    return x.mean() / stdev_s(x)


def jobson_korkie(x1, x2):
    """Jobson-Korkie test with Memmel's (2003) correction, on daily Sharpe ratios of excess returns x1 and x2."""
    sr1, sr2 = sharpe_daily(x1), sharpe_daily(x2)
    rho = np.corrcoef(x1, x2)[0, 1]
    T = len(x1)
    se = math.sqrt((2 - 2 * rho + 0.5 * (sr1 ** 2 + sr2 ** 2 - 2 * sr1 * sr2 * rho ** 2)) / T)
    z = (sr1 - sr2) / se
    return SimpleNamespace(rho=rho, se=se, z=z, p=norm_2t(z))


def drawdown(da, k):
    """Maximum drawdown, the peak before it (first day the pre-trough maximum was reached), the trough, and the
    first later day the value got back to that peak ('Not yet' if never)."""
    v, dd, dates = da["value_" + k].values, da["dd_" + k].values, da.date
    t = int(np.argmin(dd))                              # first day of the lowest drawdown
    peak_value = v[:t + 1].max()
    p = int(np.argmax(v == peak_value))
    after = np.nonzero(v[t + 1:] >= peak_value)[0]
    recovered = to_date(dates.iat[t + 1 + after[0]]) if len(after) else "Not yet"
    return SimpleNamespace(max_dd=dd.min(), peak=to_date(dates.iat[p]), trough=to_date(dates.iat[t]),
                           recovered=recovered)


# ---------------------------------------------------------------------------------------------------------------------
# 2.6 A results sheet: printed as tables, and kept cell by cell for --check
# ---------------------------------------------------------------------------------------------------------------------
class Report:
    """One results sheet of the workbook. Each row() call records the values under their workbook cell addresses
    (for --check) and adds a line to the printed table. The label is the text in column B of that row."""

    def __init__(self, sheet, title, compare="formulas"):
        self.sheet, self.title = sheet, title
        self.cells = {}           # "D37" -> value
        self.labels = {}          # row -> text expected in column B (checked before comparing that row)
        self.headers = {}         # "W5" -> text expected in that header cell (workings sheets)
        self.panels = []          # (heading, [(column letter, header)], [lines])
        # --check compares every formula cell of the sheet ("formulas") or only the cells listed here ("listed")
        self.compare = compare

    def panel(self, heading, columns):
        self.panels.append((heading, list(columns.items()), []))

    def row(self, r, label, fmt, check_label=True, **values):
        for c, v in values.items():
            self.cells[f"{c}{r}"] = v
        if check_label and label is not None:
            self.labels[r] = label
        self.panels[-1][2].append(("row", label, fmt, values))

    def text(self, cell, value):
        """A text cell compared by --check but not printed as a table row (column headers, short descriptions)."""
        self.cells[cell] = value

    def note(self, line):
        self.panels[-1][2].append(("note", line, None, None))


def fmt_value(v, fmt):
    if isinstance(v, str):
        return v
    if isinstance(v, (dt.date, pd.Timestamp, np.datetime64)):
        return to_date(v).isoformat()
    if v is None or (isinstance(v, float) and np.isnan(v)):
        return ""
    if fmt == "pct":
        return f"{v * 100:.2f}%"
    if fmt == "pct3":
        return f"{v * 100:.3f}%"
    if fmt == "year":
        return f"{v:.0f}"
    if fmt == "int":
        return f"{v:,.0f}"
    if fmt == "p":
        return f"{v:.4f}" if v >= 0.0001 or v == 0 else f"{v:.2e}"
    if fmt == "usd":
        return f"{v:,.2f}"
    if fmt == "e":
        return f"{v:.2e}"
    if fmt == "num2":
        return f"{v:,.0f}" if float(v).is_integer() else f"{v:,.2f}"
    return f"{v:.4f}"


def pick_fmt(fmt, c):
    """A row's number format: one code for every column, or a dict {column: code}."""
    return fmt.get(c, "") if isinstance(fmt, dict) else fmt


def print_report(rep, width=140):
    print("\n" + "=" * width)
    print(f"{rep.sheet.upper()}  |  {rep.title}")
    print("=" * width)
    label_w = 66
    for heading, columns, lines in rep.panels:
        print(f"\n--- {heading} ---")
        rows = [ln for ln in lines if ln[0] == "row"]
        colw = {}
        for c, h in columns:
            vals = [fmt_value(ln[3][c], pick_fmt(ln[2], c)) for ln in rows if c in ln[3]]
            vals = [v for v in vals if len(v) <= 40]
            colw[c] = max([len(h)] + [len(v) for v in vals] + [10]) + 2
        if columns:
            print(" " * label_w + "".join(h.rjust(colw[c]) for c, h in columns))
        for kind, label, fmt, values in lines:
            if kind == "note":
                print("  " + label)
                continue
            lab = str(label if label is not None else "")
            lab = lab if len(lab) <= label_w - 2 else lab[:label_w - 3] + "…"
            cells = []
            long_texts = []
            for c, h in columns:
                if c in values:
                    s = fmt_value(values[c], pick_fmt(fmt, c))
                    if len(s) > 40:
                        long_texts.append(s)
                        s = "(see below)"
                    cells.append(s.rjust(colw[c]))
                else:
                    cells.append(" " * colw[c])
            print(lab.ljust(label_w) + "".join(cells))
            for s in long_texts:
                print("    " + s)

# ---------------------------------------------------------------------------------------------------------------------
# 2.7 Settings sheet
# ---------------------------------------------------------------------------------------------------------------------
def report_settings(da):
    rep = Report("Settings", "every assumption used by the analysis (named constants at the top of this script)",
                 compare="listed")
    rep.panel("Inputs and derived constants", {"C": "Value"})
    rows = [
        (7, "Significance level (two-sided)", "ALPHA", ALPHA, "num"),
        (8, "Critical value for that level (normal)", "ZCRIT", ZCRIT, "num"),
        (9, "Power for the minimum detectable effect", "POWER", POWER, "num"),
        (10, "z for that power (normal)", "ZPOW", ZPOW, "num"),
        (11, "MDE multiplier", "MDE_MULT", MDE_MULT, "num"),
        (12, "Smallest effect worth detecting, per trade", "MIN_EFFECT", MIN_EFFECT, "num"),
        (13, "Trading days per year", "DAYS_PER_YEAR", DAYS_PER_YEAR, "int"),
        (14, "Minimum days needed to estimate a beta", "MIN_BETA_DAYS", MIN_BETA_DAYS, "int"),
        (15, "Newey–West lags", "NW_LAGS", NW_LAGS, "int"),
        (16, "Last year of the first half", "SPLIT_YEAR", SPLIT_YEAR, "year"),
        (17, "Winsorising: lower percentile (V3)", "WINSOR_LO", WINSOR_LO, "num"),
        (18, "Winsorising: upper percentile (V3)", "WINSOR_HI", WINSOR_HI, "num"),
        (19, "Trades dropped at each end (V6)", "TRIM_N", TRIM_N, "int"),
        (20, "Year left out in V4 (first)", "EXCL_YEAR_1", EXCL_YEAR_1, "year"),
        (21, "Year left out in V4 (second)", "EXCL_YEAR_2", EXCL_YEAR_2, "year"),
        (22, "Last day with Fama-French factors", "FF_LAST (from the data)", to_date(da.attrs["ff_last"]), "num"),
        (23, "Start of the 'from 2020' check", "START_2020", START_2020, "num"),
        (24, "Break test: share of dates kept out at each end", "BREAK_TRIM", BREAK_TRIM, "num"),
        (25, "Break test: critical value at 5%", "BAI_PERRON_CV_5", BAI_PERRON_CV_5, "num"),
        (26, "Break test: critical value at 1%", "BAI_PERRON_CV_1", BAI_PERRON_CV_1, "num"),
        (27, "Largest quarterly trades left out (E5)", "N_DROP_V1", N_DROP_V1, "int"),
    ]
    for r, label, name, value, fmt in rows:
        rep.row(r, f"{label}  [{name}]", fmt, check_label=False, C=value)
        rep.labels[r] = label
    rep.note(f"Worked example in Tests 1 and 2: {WORKED_EXAMPLE}. Beta window: days -{BETA_FIRST_DAY} to "
             f"-{BETA_LAST_DAY}. The account starts on the first trading day on or after {text_dmy(ACCOUNT_START)}.")
    return rep


# ---------------------------------------------------------------------------------------------------------------------
# 2.8 Workings summaries (Trade Returns, Beta Estimation, Holding Days, Daily Account)
# ---------------------------------------------------------------------------------------------------------------------
def print_workings(raw, tr, beta, hd, da):
    print("\n" + "=" * 140)
    print("WORKINGS  |  Trade Returns, Beta Estimation, Holding Days, Daily Account "
          "(--save-workings writes them out in full)")
    print("=" * 140)
    print(f"Workbook read: {raw.path}")
    print(f"Trades: {raw.n} additions (rows {FIRST_ROW}-{FIRST_ROW + raw.n - 1}), announced "
          f"{tr.announced.min():%d %b %Y} to {tr.announced.max():%d %b %Y}, on {tr.announced.nunique()} dates; "
          f"{(tr.quarterly == 'Yes').sum()} at a quarterly rebalance.")
    print(f"Holding days: {len(hd):,} position-days (median {tr.days.median():.0f} days, range "
          f"{tr.days.min()}-{tr.days.max()}).")
    few = beta.days.values < MIN_BETA_DAYS
    print(f"Beta: {(~few).sum()} estimated, {few.sum()} set to 1 "
          f"({', '.join(f'{n} ({d} days)' for n, d in zip(tr.name.values[few], beta.days.values[few]))}).")
    print(f"Daily Account: {len(da):,} trading days, {to_date(da.date.iat[0])} to {to_date(da.date.iat[-1])}; "
          f"Fama-French factors to {to_date(da.attrs['ff_last'])}.")
    show = tr[["row", "name", "announced", "included", "days", "AR_A", "AR_B", "beta", "AR_A_adj", "AR_B_adj", "V5"]]
    print("\nTrade Returns, the 5 largest and 5 smallest Window A abnormal returns:")
    pick = pd.concat([show.nlargest(5, "AR_A"), show.nsmallest(5, "AR_A")])
    with pd.option_context("display.width", 160, "display.max_columns", 20, "display.float_format", "{:.4f}".format):
        print(pick.to_string(index=False))


def save_workings(folder, tr, beta, hd, da):
    folder = Path(folder)
    folder.mkdir(parents=True, exist_ok=True)
    tr.assign(beta_days=beta.days.values, beta_slope=beta.slope.values).to_csv(folder / "trade_returns.csv",
                                                                               index=False)
    hd.to_csv(folder / "holding_days.csv", index=False)
    da.to_csv(folder / "daily_account.csv", index=False)
    print(f"\nWorkings saved in {folder}: trade_returns.csv, holding_days.csv, daily_account.csv")


def workings_reports(tr, beta, hd, da):
    """The workings sheets, column by column, for --check (not printed: --save-workings writes them out)."""
    out = []
    rep = Report("Trade Returns", "", compare="listed")
    cols = {"E": "announced", "F": "included", "G": "year", "I": "s_ann_close", "J": "s_buy_open",
            "K": "s_ann1_close", "L": "s_pre_close", "M": "s_inc_open", "N": "s_inc_close", "O": "m_ann_close",
            "P": "m_buy_open", "Q": "m_ann1_close", "R": "m_pre_close", "S": "m_inc_open", "T": "m_inc_close",
            "U": "R_A", "V": "M_A", "W": "AR_A", "X": "R_B", "Y": "M_B", "Z": "AR_B", "AA": "beta",
            "AB": "AR_A_adj", "AC": "AR_B_adj", "AD": "V5_stock", "AE": "V5_spy", "AF": "V5",
            "AG": "CAR_ann_close", "AH": "CAR_buy_open", "AI": "CAR_ann1_close", "AJ": "CAR_pre_close",
            "AK": "CAR_inc_open", "AL": "CAR_inc_close", "AM": "days", "AN": "gap_stock_A", "AO": "gap_stock_B",
            "AP": "gap_spy_A", "AQ": "gap_spy_B", "AR": "label", "B": "ref", "C": "name"}
    for c, k in cols.items():
        for i, v in enumerate(tr[k].values):
            rep.cells[f"{c}{FIRST_ROW + i}"] = v
    rep.headers = {"W5": "Abnormal return AR_A", "Z5": "Abnormal return AR_B", "AA5": "Beta", "AF5": "Abnormal return",
                   "AK5": "At inclusion open"}
    out.append(rep)
    rep = Report("Beta Estimation", "", compare="listed")
    for i in range(len(beta)):
        rep.cells[f"D{FIRST_ROW + i}"] = beta.days.iat[i]
        rep.cells[f"F{FIRST_ROW + i}"] = beta.slope.iat[i]          # blank (NaN) = "too few days" in the sheet
        rep.cells[f"G{FIRST_ROW + i}"] = beta.beta.iat[i]
    rep.headers = {"D5": "Days with a return (−250 to −30)", "G5": "Beta used"}
    out.append(rep)
    rep = Report("Holding Days", "", compare="listed")
    cols = {"B": "row", "D": "day", "F": "date", "G": "stock_open", "H": "stock_close", "I": "stock_prev_close",
            "J": "spy_open", "K": "spy_close", "L": "spy_prev_close", "N": "stock_A", "O": "spy_A", "P": "ar_A",
            "Q": "stock_B", "R": "spy_B", "S": "ar_B", "T": "long_spy_A", "U": "long_spy_B"}
    last = hd.groupby("trade").day.transform("max").values
    for c, k in cols.items():
        for j, v in enumerate(hd[k].values):
            rep.cells[f"{c}{FIRST_ROW + j}"] = v
    for j, v in enumerate(last):
        rep.cells[f"E{FIRST_ROW + j}"] = v
    rep.headers = {"D5": "Holding day (1 = Ann+1)", "F5": "Date", "T5": "Window A", "U5": "Window B"}
    out.append(rep)
    rep = Report("Daily Account", "", compare="listed")
    cols = {"B": "date", "C": "rf", "D": "positions", "E": "spy_close", "F": "spy_prev_close", "G": "B1",
            "I": "avg_A", "J": "A", "K": "avg_B", "L": "B", "M": "x_A", "N": "x_B1", "O": "x_B",
            "P": "value_A", "Q": "value_B1", "R": "value_B", "S": "dd_A", "T": "dd_B1", "U": "dd_B"}
    for c, k in cols.items():
        for j, v in enumerate(da[k].values):
            rep.cells[f"{c}{FIRST_ROW + j}"] = v
    rep.headers = {"B5": "Date", "G5": "B1 return", "J5": "Long + SPY", "P5": "Long + SPY, A",
                   "Q5": "B1 SPY buy and hold"}
    out.append(rep)
    return out


# ---------------------------------------------------------------------------------------------------------------------
# 2.9 Test 1 Per Trade and Test 2 Risk-Adjusted
# ---------------------------------------------------------------------------------------------------------------------
def per_trade_sections(rep, a, b, test, off, beat_word):
    """Sections a-f and the checks of Test 1 (off = 0) or Test 2 (off = 12; checks at 116 instead of 135).
    a, b = stat blocks of Window A (column D) and Window B (column E)."""
    sig5 = f"Significant at the {text_pct(ALPHA)} level?"
    rep.panel(f"{test}a. Average abnormal return, each trade once", {"D": "Window A", "E": "Window B"})
    rep.row(36 + off, "N: number of trades", "int", D=a.N, E=b.N)
    rep.row(37 + off, "Mean abnormal return", "pct", D=a.mean, E=b.mean)
    rep.row(38 + off, "s: standard deviation", "pct", D=a.s, E=b.s)
    rep.row(39 + off, "SE: standard error of the mean", "pct", D=a.se, E=b.se)
    rep.row(40 + off, "t-statistic", "num", D=a.t, E=b.t)
    rep.row(41 + off, "p-value (two-sided)", "p", D=a.p, E=b.p)
    rep.row(42 + off, sig5, "", D=sig_word(a.p, a.mean), E=sig_word(b.p, b.mean))
    rep.panel(f"{test}b. Clustered by announcement date (each date counts once)", {"D": "Window A", "E": "Window B"})
    rep.row(50 + off, "G: number of announcement dates", "int", D=a.G, E=b.G)
    rep.row(51 + off, "…of which dates with 2 or more trades", "int", D=a.multi, E=b.multi)
    rep.row(52 + off, "Mean of the date averages", "pct", D=a.meanG, E=b.meanG)
    rep.row(53 + off, "s of the date averages", "pct", D=a.sdG, E=b.sdG)
    rep.row(54 + off, "SE (clustered)", "pct", D=a.seG, E=b.seG)
    rep.row(55 + off, "t-statistic (clustered)", "num", D=a.tG, E=b.tG)
    rep.row(56 + off, "p-value (two-sided)", "p", D=a.pG, E=b.pG)
    rep.row(57 + off, sig5, "", D=sig_word(a.pG, a.meanG), E=sig_word(b.pG, b.meanG))
    rep.panel(f"{test}c. {text_pct(1 - ALPHA)} confidence interval for the mean of the date averages",
              {"D": "Window A", "E": "Window B"})
    rep.row(62 + off, f"t*: the t value that leaves {text_pct(ALPHA / 2, 1)} in each tail", "num",
            D=a.t_star, E=b.t_star)
    rep.text(f"C{62 + off}", f"T.INV.2T({text_pct(ALPHA)}, G − 1)")
    rep.row(63 + off, "Lower end", "pct", D=a.lo, E=b.lo)
    rep.row(64 + off, "Upper end", "pct", D=a.hi, E=b.hi)
    rep.row(65 + off, "Does the interval exclude zero?", "", D=yes(a.lo > 0 or a.hi < 0), E=yes(b.lo > 0 or b.hi < 0))
    rep.panel(f"{test}d. Median and Wilcoxon signed-rank test", {"D": "Window A", "E": "Window B"})
    rep.row(72 + off, "Median abnormal return", "pct", D=a.median, E=b.median)
    rep.row(73 + off, "n: trades with a non-zero return", "int", D=a.n, E=b.n)
    rep.row(74 + off, "W+: sum of the ranks of the positive trades", "num2", D=a.W, E=b.W)
    rep.row(75 + off, "Expected W+ if the strategy does nothing", "num2", D=a.EW, E=b.EW)
    rep.row(76 + off, "Variance of W+ if the strategy does nothing", "num2", D=a.VW, E=b.VW)
    rep.row(77 + off, "z-statistic", "num", D=a.z, E=b.z)
    rep.row(78 + off, "p-value (two-sided)", "p", D=a.pW, E=b.pW)
    rep.row(79 + off, sig5, "", D=sig_word(a.pW, a.z), E=sig_word(b.pW, b.z))
    rep.panel(f"{test}e. Sign test", {"D": "Window A", "E": "Window B"})
    beat50 = lambda s: ("Yes, above 50%" if s.beat > 0.5 else "Yes, below 50%") if s.pS < ALPHA else "No"
    rep.row(84 + off, f"k: trades that beat {beat_word}", "int", D=a.k, E=b.k)
    rep.row(85 + off, "n: trades with a non-zero return", "int", D=a.n, E=b.n)
    rep.row(86 + off, f"% of trades that beat {beat_word}", "pct", D=a.beat, E=b.beat)
    rep.row(87 + off, "p-value (two-sided sign test)", "p", D=a.pS, E=b.pS)
    rep.row(88 + off, "Significantly different from 50%?", "", D=beat50(a), E=beat50(b))
    rep.panel(f"{test}f. Minimum detectable effect (power {text_pct(POWER)})", {"D": "Window A", "E": "Window B"})
    rep.row(93 + off, f"SE (clustered, from {test}b)", "pct", D=a.seG, E=b.seG)
    rep.row(94 + off, "MDE: minimum detectable effect", "pct", D=a.mde, E=b.mde)
    rep.row(95 + off, "Smallest effect worth detecting", "pct", D=MIN_EFFECT, E=MIN_EFFECT)
    rep.row(96 + off, "Could the test detect it?", "", D=yes(a.mde <= MIN_EFFECT), E=yes(b.mde <= MIN_EFFECT))


def verdict_checks(rep, a, b, r0, test):
    """The six checks behind Test 1's rule, and the verdict, rows r0..r0+6."""
    rep.panel("Putting it together (Test 1's rule)", {"D": "Window A", "E": "Window B"})
    rep.row(r0, "1. Average above zero?", "", D=yes(a.meanG > 0), E=yes(b.meanG > 0))
    rep.row(r0 + 1, "2. …and significantly so?", "", D=yes(a.pG < ALPHA and a.meanG > 0),
            E=yes(b.pG < ALPHA and b.meanG > 0))
    rep.text(f"C{r0 + 1}", f"{test}b: clustered p below {text_pct(ALPHA)} with the mean above 0")
    rep.row(r0 + 2, f"3. {text_pct(1 - ALPHA)} interval entirely above zero?", "", D=yes(a.lo > 0), E=yes(b.lo > 0))
    rep.row(r0 + 3, "4. Median above zero?", "", D=yes(a.median > 0), E=yes(b.median > 0))
    rep.row(r0 + 4, "5. …and significantly so?", "", D=yes(a.pW < ALPHA and a.z > 0), E=yes(b.pW < ALPHA and b.z > 0))
    rep.text(f"C{r0 + 4}", f"{test}d: Wilcoxon p below {text_pct(ALPHA)} with z above 0")
    rep.row(r0 + 5, f"6. Could the test detect {text_pct(MIN_EFFECT)} per trade?", "",
            D=yes(a.mde <= MIN_EFFECT), E=yes(b.mde <= MIN_EFFECT))
    rep.row(r0 + 6, "Result for this window", "", D=verdict(a), E=verdict(b))


def worked_example(rep, tr, r0, risk_adjusted=False):
    """The worked example (one trade, both windows), rows r0..r0+9 (and +10, +11 for the beta lines in Test 2)."""
    i = int(np.nonzero(tr.label.values == WORKED_EXAMPLE)[0][0])
    t = tr.iloc[i]
    rep.panel(f"Worked example: {WORKED_EXAMPLE}", {"D": "Window A", "E": "Window B"})
    rep.row(r0, "Row in the data sheets", "int", D=int(t.row))
    rep.row(r0 + 1, "Announced", "", D=to_date(t.announced))
    rep.row(r0 + 2, "Included (first day in the index)", "", D=to_date(t.included))
    rep.row(r0 + 3, "Stock: buy at the Ann+1 open ($)", "usd", D=t.s_buy_open, E=t.s_buy_open)
    rep.row(r0 + 4, "Stock: sell ($)", "usd", D=t.s_inc_open, E=t.s_inc_close)
    rep.row(r0 + 5, "Stock return R", "pct", D=t.R_A, E=t.R_B)
    rep.row(r0 + 6, "SPY: buy at the Ann+1 open ($)", "usd", D=t.m_buy_open, E=t.m_buy_open)
    rep.row(r0 + 7, "SPY: sell ($)", "usd", D=t.m_inc_open, E=t.m_inc_close)
    rep.row(r0 + 8, "SPY return M", "pct", D=t.M_A, E=t.M_B)
    rep.row(r0 + 9, "Abnormal return AR = R − M", "pct", D=t.AR_A, E=t.AR_B)
    if risk_adjusted:
        rep.row(r0 + 10, "Beta (Beta Estimation sheet)", "num", D=t.beta)
        rep.row(r0 + 11, "Risk-adjusted AR = R − β × M", "pct", D=t.AR_A_adj, E=t.AR_B_adj)


def report_test1(tr):
    rep = Report("Test 1 Per Trade",
                 "does buying at the Ann+1 open and selling at inclusion beat SPY over the same hours?")
    dates = tr.announced.values
    a, b = stat_block(tr.AR_A, dates), stat_block(tr.AR_B, dates)
    rep.panel("Answer", {"D": "Result"})
    rep.row(9, "Test 1 result", "", D=combine(verdict(a), verdict(b)))
    worked_example(rep, tr, 22)
    per_trade_sections(rep, a, b, "1", 0, "SPY")
    # Figure 1: average abnormal return since the announcement close, at six moments
    rep.panel("Figure 1. Average abnormal return since the announcement close (each date once; band = mean ± 1.96 SE)",
              {"D": "Mean (dates)", "E": "SE", "F": f"Lower {text_pct(1 - ALPHA)}", "G": f"Upper {text_pct(1 - ALPHA)}",
               "H": "Mean (trades)"})
    rep.text("F102", f"Lower {text_pct(1 - ALPHA)}")
    rep.text("G102", f"Upper {text_pct(1 - ALPHA)}")
    moments = [(103, "Announcement close", "CAR_ann_close"), (104, "Ann+1 open (we buy)", "CAR_buy_open"),
               (105, "Ann+1 close", "CAR_ann1_close"), (106, "Close the day before inclusion", "CAR_pre_close"),
               (107, "Inclusion open (sell, Window A)", "CAR_inc_open"),
               (108, "Inclusion close (sell, Window B)", "CAR_inc_close")]
    for r, label, c in moments:
        by_date = tr.groupby("announced")[c].mean()
        m, se = by_date.mean(), stdev_s(by_date.values) / math.sqrt(a.G)
        rep.row(r, label, "pct", D=m, E=se, F=m - ZCRIT * se, G=m + ZCRIT * se, H=tr[c].mean())
    verdict_checks(rep, a, b, 135, "1")
    return rep, a, b


def report_test2(tr, beta, t1a, t1b):
    rep = Report("Test 2 Risk-Adjusted", "the same tests on R − β × M (beta from days −250 to −30)")
    dates = tr.announced.values
    a, b = stat_block(tr.AR_A_adj, dates), stat_block(tr.AR_B_adj, dates)
    va, vb = verdict(a), verdict(b)
    same_a, same_b = yes(verdict(t1a) == va), yes(verdict(t1b) == vb)
    rep.panel("Answer", {"D": "Result"})
    rep.row(9, "Test 2 result", "", D=combine(va, vb))
    rep.row(10, "Same answer as Test 1?", "", D=yes(same_a == "Yes" and same_b == "Yes"))
    rep.panel("Betas", {"D": "Value"})
    used = beta.beta.values
    rep.row(20, "Trades with an estimated beta", "int", D=int((beta.days >= MIN_BETA_DAYS).sum()))
    rep.row(21, "Trades with beta set to 1", "int", D=int((beta.days < MIN_BETA_DAYS).sum()))
    rep.row(22, "Average beta used", "num", D=used.mean())
    rep.row(23, "Median beta used", "num", D=float(np.median(used)))
    rep.row(24, "Lowest beta used", "num", D=used.min())
    rep.row(25, "Highest beta used", "num", D=used.max())
    rep.row(26, "Share of stocks with a beta above 1", "pct", D=(used > 1).sum() / len(used))
    worked_example(rep, tr, 32, risk_adjusted=True)
    per_trade_sections(rep, a, b, "2", 12, "β × SPY")
    verdict_checks(rep, a, b, 116, "2")
    rep.panel("Compared with Test 1", {"D": "Window A", "E": "Window B"})
    rep.row(126, "Average, each date once: Test 1", "pct", D=t1a.meanG, E=t1b.meanG)
    rep.row(127, "Average, each date once: risk-adjusted", "pct", D=a.meanG, E=b.meanG)
    rep.row(128, "Difference", "pct", D=a.meanG - t1a.meanG, E=b.meanG - t1b.meanG)
    rep.row(129, "Average, each trade once: Test 1", "pct", D=t1a.mean, E=t1b.mean)
    rep.row(130, "Average, each trade once: risk-adjusted", "pct", D=a.mean, E=b.mean)
    rep.row(131, "Result, Test 1", "", D=verdict(t1a), E=verdict(t1b))
    rep.row(132, "Result, risk-adjusted", "", D=va, E=vb)
    rep.row(133, "Same answer as Test 1?", "", D=same_a, E=same_b)
    return rep


# ---------------------------------------------------------------------------------------------------------------------
# 2.10 Test 3 Portfolio
# ---------------------------------------------------------------------------------------------------------------------
def account_panel(rep, da, k, off, regs, z_last):
    """Panel A (k = 'A', off = 0) or Panel B (k = 'B', off = 42): the Long + SPY account (column D) against B1 (E)."""
    T = len(da)
    res = {}
    for col, key in [("D", k), ("E", "B1")]:
        r, x, v = da[key].values, da["x_" + key].values, da["value_" + key].values
        res[col] = SimpleNamespace(mean=r.mean(), ann=v[-1] ** (DAYS_PER_YEAR / T) - 1, end=v[-1],
                                   vol=stdev_s(r) * math.sqrt(DAYS_PER_YEAR), xmean=x.mean(), xsd=stdev_s(x),
                                   srd=x.mean() / stdev_s(x), dd=drawdown(da, key))
    D, E = res["D"], res["E"]
    jk = jobson_korkie(da["x_" + k].values, da.x_B1.values)
    capm, ff4 = regs
    name = "Long + SPY, Window " + k
    rep.panel(f"Panel {k}: {name} (D) against B1, SPY bought and held (E)", {"D": name, "E": "B1"})
    rep.row(34 + off, "Days (T)", "int", D=T, E=T)
    rep.row(35 + off, "% of days holding added stocks", "pct", D=(da.positions > 0).sum() / T)
    rep.row(36 + off, "Average daily return", "pct3", D=D.mean, E=E.mean)
    rep.row(37 + off, "Annualised return", "pct", D=D.ann, E=E.ann)
    rep.text(f"C{37 + off}", f"(value of $1 at the end)^({DAYS_PER_YEAR} ÷ T) − 1")
    rep.row(38 + off, "Return per year above B1", "pct", D=D.ann - E.ann)
    rep.row(39 + off, "Value of $1 at the end", "usd", D=D.end, E=E.end)
    rep.row(40 + off, "Annualised volatility", "pct", D=D.vol, E=E.vol)
    rep.text(f"C{40 + off}", f"STDEV.S of the daily returns × √{DAYS_PER_YEAR}")
    rep.row(41 + off, "Average daily excess return", "pct3", D=D.xmean, E=E.xmean)
    rep.row(42 + off, "St. dev. of daily excess return", "pct3", D=D.xsd, E=E.xsd)
    rep.row(43 + off, "Sharpe ratio, daily", "num", D=D.srd, E=E.srd)
    rep.row(44 + off, "Sharpe ratio, annualised", "num", D=D.srd * math.sqrt(DAYS_PER_YEAR),
            E=E.srd * math.sqrt(DAYS_PER_YEAR))
    rep.text(f"C{44 + off}", f"Daily Sharpe × √{DAYS_PER_YEAR}")
    rep.row(46 + off, "ρ: correlation with B1", "num", D=jk.rho)
    rep.row(47 + off, "SE of the Sharpe difference", "num", D=jk.se)
    rep.row(48 + off, "z-statistic", "num", D=jk.z)
    rep.row(49 + off, "p-value (two-sided)", "p", D=jk.p)
    above, below = yes(jk.z > 0 and jk.p < ALPHA), yes(jk.z < 0 and jk.p < ALPHA)
    rep.row(50 + off, "Significantly above B1?", "", D=above)
    rep.row(51 + off, "Significantly below B1?", "", D=below)
    rep.row(54 + off, "Days in the regressions", "int", D=z_last)
    rep.row(55 + off, "CAPM alpha, per year", "pct", D=capm.alpha_year)
    rep.text(f"C{55 + off}", f"Daily α × {DAYS_PER_YEAR}")
    rep.row(56 + off, "t-statistic (Newey–West)", "num", D=capm.t)
    rep.row(57 + off, "p-value (two-sided)", "p", D=t_dist_2t(capm.t, z_last - 2))
    rep.row(58 + off, "CAPM beta", "num", D=capm.b[1])
    rep.row(59 + off, "Four-factor alpha, per year", "pct", D=ff4.alpha_year)
    rep.text(f"C{59 + off}", f"Daily α × {DAYS_PER_YEAR}, with Mkt-RF, SMB, HML and Mom")
    rep.row(60 + off, "t-statistic (Newey–West)", "num", D=ff4.t)
    rep.row(61 + off, "p-value (two-sided)", "p", D=t_dist_2t(ff4.t, z_last - 5))
    rep.row(62 + off, "Loading on Mkt-RF", "num", D=ff4.b[1])
    rep.row(63 + off, "Loading on SMB (small stocks)", "num", D=ff4.b[2])
    rep.row(64 + off, "Loading on HML (value stocks)", "num", D=ff4.b[3])
    rep.row(65 + off, "Loading on Mom (recent winners)", "num", D=ff4.b[4])
    rep.row(68 + off, "Maximum drawdown", "pct", D=D.dd.max_dd, E=E.dd.max_dd)
    rep.row(69 + off, "Peak before the fall", "", D=D.dd.peak, E=E.dd.peak)
    rep.row(70 + off, "Trough", "", D=D.dd.trough, E=E.dd.trough)
    rep.row(71 + off, "Recovered by", "", D=D.dd.recovered, E=E.dd.recovered)
    return SimpleNamespace(sharpe_above=above, capm=yes(capm.alpha_year > 0 and capm.t > ZCRIT),
                           ff4=yes(ff4.alpha_year > 0 and ff4.t > ZCRIT))


def report_test3(tr, hd, da, regs):
    rep = Report("Test 3 Portfolio", "the strategy run as an account (Long + SPY) against SPY bought and held (B1)")
    T = len(da)
    z_last = regs["full_T"]
    no_pos = da.positions.values == 0
    rep.panel("Checks on the workings", {"D": "Value", "E": "OK?"})
    rep.row(24, "Position-days in the Holding Days sheet", "int", D=len(hd), E=yes(len(hd) == int(tr.days.sum())))
    rep.row(25, f"Trading days held, summed over the {len(tr)} trades", "int", D=int(tr.days.sum()))
    rep.row(26, "Positions held, summed over every day of the account", "int", D=int(da.positions.sum()),
            E=yes(int(da.positions.sum()) == len(hd)))
    g_stock = np.abs(tr[["gap_stock_A", "gap_stock_B"]].values).max()
    g_spy = np.abs(tr[["gap_spy_A", "gap_spy_B"]].values).max()
    rep.row(27, "Largest gap between a stock's chained pieces and its trade return", "e", D=g_stock,
            E=yes(g_stock < 1e-6))
    rep.row(28, "Largest gap between SPY's chained pieces and its window return", "e", D=g_spy, E=yes(g_spy < 1e-6))
    bad = int((no_pos & (np.abs(da.A.values - da.B1.values) > 1e-7)).sum()
              + (no_pos & (np.abs(da.B.values - da.B1.values) > 1e-7)).sum())
    rep.row(29, "Long + SPY equals B1 on every day with no position", "int", D=bad, E=yes(bad == 0))
    found = int(da.price_found.sum())
    rep.row(30, "Days whose SPY price matches the date", "", D=f"{found:,} of {T:,}", E=yes(found == T))
    A = account_panel(rep, da, "A", 0, (regs["capm_A"], regs["ff4_A"]), z_last)
    B = account_panel(rep, da, "B", 42, (regs["capm_B"], regs["ff4_B"]), z_last)
    rep.panel("Putting it together", {"D": "Window A", "E": "Window B"})
    rep.row(120, "Sharpe ratio reliably above B1?", "", D=A.sharpe_above, E=B.sharpe_above)
    rep.row(121, f"CAPM alpha above zero with t > {ZCRIT:.2f}?", "", D=A.capm, E=B.capm)
    rep.row(122, f"Four-factor alpha above zero with t > {ZCRIT:.2f}?", "", D=A.ff4, E=B.ff4)
    res = {c: ("Works" if (w.sharpe_above, w.capm, w.ff4) == ("Yes",) * 3 else "Doesn't work") for c, w in
           [("D", A), ("E", B)]}
    rep.row(123, "Result for this window", "", D=res["D"], E=res["E"])
    rep.panel("Answer", {"D": "Result"})
    rep.row(8, "Test 3 result", "", D="Works" if "Works" in res.values() else "Doesn't work")
    rep.panels.insert(0, rep.panels.pop())
    return rep, res


# ---------------------------------------------------------------------------------------------------------------------
# 2.11 Test 4 Robustness
# ---------------------------------------------------------------------------------------------------------------------
STAT_ROWS = ["N: trades", "Mean abnormal return (each trade once)", "Median", "% of trades that beat SPY",
             "Announcement dates (G)", "Mean of the date averages (each date once)", "Clustered SE",
             "Clustered t = mean of date averages ÷ SE", "p-value: T.DIST.2T(|t|, G − 1)",
             None, "Wilcoxon n (non-zero trades)", "W+: sum of ranks of positive trades", "Expected W+ = n(n + 1) ÷ 4",
             "z = (W+ − expected) ÷ √[n(n + 1)(2n + 1) ÷ 24]", "Wilcoxon p-value", None,
             "Average significant? (clustered t-test)", "Verdict (Test 1's rule)", "Same verdict as Test 1?"]


def robust_block(rep, r0, heading, columns, test1_verdicts):
    """Test 4's statistic block (rows r0..r0+18). columns = {letter: (header, x, mask, dates)}; the statistics use
    only the trades in mask, with date averages and Wilcoxon ranks computed within that subset."""
    blocks = {c: stat_block(x[m], dates[m]) for c, (_, x, m, dates) in columns.items()}
    rep.panel(heading, {c: v[0] for c, v in columns.items()})
    labels = list(STAT_ROWS)
    labels[9] = f"Lower end of the {text_pct(1 - ALPHA)} interval: mean − t* × SE"
    labels[15] = f"MDE = ({ZCRIT:.2f} + {ZPOW:.2f}) × clustered SE"
    fields = [("N", "int"), ("mean", "pct"), ("median", "pct"), ("beat", "pct"), ("G", "int"), ("meanG", "pct"),
              ("seG", "pct"), ("tG", "num"), ("pG", "p"), ("lo", "pct"), ("n", "int"), ("W", "num2"),
              ("EW", "num2"), ("z", "num"), ("pW", "p"), ("mde", "pct")]
    for j, (f, fmt) in enumerate(fields):
        rep.row(r0 + j, labels[j], fmt, **{c: getattr(b, f) for c, b in blocks.items()})
    rep.row(r0 + 16, labels[16], "", **{c: sig_word(b.pG, b.meanG) for c, b in blocks.items()})
    rep.row(r0 + 17, labels[17], "", **{c: verdict(b) for c, b in blocks.items()})
    rep.row(r0 + 18, labels[18], "", **{c: yes(verdict(b) == test1_verdicts[c]) for c, b in blocks.items()})
    return blocks


def report_test4(tr, da, regs, t1a, t1b):
    rep = Report("Test 4 Robustness", "does the Test 1 answer hold in each half and in variants of the sample?")
    d = tr.announced.values
    xA, xB, xV5 = tr.AR_A.values, tr.AR_B.values, tr.V5.values
    year = tr.year.values
    vA, vB = verdict(t1a), verdict(t1b)
    y0, y1 = int(year.min()), int(year.max())
    H1, H2 = year <= SPLIT_YEAR, year > SPLIT_YEAR
    ALL = np.ones(len(tr), bool)
    Q, OFF = tr.quarterly.values == "Yes", tr.quarterly.values == "No"
    V4 = ~np.isin(year, [EXCL_YEAR_1, EXCL_YEAR_2])
    # V3: winsorise at the 1st and 99th percentiles (PERCENTILE.INC = numpy's default linear interpolation)
    loA, hiA = np.percentile(xA, 100 * WINSOR_LO), np.percentile(xA, 100 * WINSOR_HI)
    loB, hiB = np.percentile(xB, 100 * WINSOR_LO), np.percentile(xB, 100 * WINSOR_HI)
    wA, wB = np.minimum(np.maximum(xA, loA), hiA), np.minimum(np.maximum(xB, loB), hiB)
    # V6: drop the TRIM_N largest and TRIM_N smallest (RANK.AVG from the top and from the bottom both > TRIM_N)
    trim = lambda x: (rank_avg(-x) > TRIM_N) & (rank_avg(x) > TRIM_N)
    V6A, V6B = trim(xA), trim(xB)

    # 4a: first and second half
    h = {"C": f"First half, A ({y0}–{SPLIT_YEAR})", "D": f"Second half, A ({SPLIT_YEAR + 1}–{y1})",
         "E": f"First half, B ({y0}–{SPLIT_YEAR})", "F": f"Second half, B ({SPLIT_YEAR + 1}–{y1})"}
    for c, lab in h.items():
        rep.text(f"{c}17", lab)
    halves = robust_block(rep, 18, "4a. First half and second half (by announcement year)",
                          {"C": (h["C"], xA, H1, d), "D": (h["D"], xA, H2, d), "E": (h["E"], xB, H1, d),
                           "F": (h["F"], xB, H2, d)}, {"C": vA, "D": vA, "E": vB, "F": vB})
    # Do the halves differ? Welch t-test and Mann-Whitney
    rep.panel("4a. Do the halves differ? (Welch t-test and Mann–Whitney)", {"C": "Window A", "D": "Window B"})
    res = {}
    for c, x in [("C", xA), ("D", xB)]:
        a, b = x[H1], x[H2]
        n1, n2 = len(a), len(b)
        s1, s2 = stdev_s(a), stdev_s(b)
        se = math.sqrt(s1 ** 2 / n1 + s2 ** 2 / n2)
        diff = b.mean() - a.mean()
        df = (s1 ** 2 / n1 + s2 ** 2 / n2) ** 2 / ((s1 ** 2 / n1) ** 2 / (n1 - 1) + (s2 ** 2 / n2) ** 2 / (n2 - 1))
        R1 = rank_avg(x)[H1].sum()                       # ranks among all trades of this window
        U = R1 - n1 * (n1 + 1) / 2
        EU = n1 * n2 / 2
        zU = (U - EU) / math.sqrt(n1 * n2 * (n1 + n2 + 1) / 12)
        res[c] = dict(n1=n1, n2=n2, m1=a.mean(), m2=b.mean(), s1=s1, s2=s2, diff=diff, se=se, t=diff / se, df=df,
                      p=t_dist_2t(diff / se, df), R1=R1, U=U, EU=EU, z=zU, pU=norm_2t(zU))
    rows = [(40, "n, first half", "n1", "int"), (41, "n, second half", "n2", "int"),
            (42, "Mean, first half", "m1", "pct"), (43, "Mean, second half", "m2", "pct"),
            (44, "St. dev., first half", "s1", "pct"), (45, "St. dev., second half", "s2", "pct"),
            (46, "Difference (second − first)", "diff", "pct"),
            (47, "SE of the difference = √(s₁²/n₁ + s₂²/n₂)", "se", "pct"),
            (48, "Welch t = difference ÷ SE", "t", "num"), (49, "Welch degrees of freedom", "df", "num2"),
            (50, "Welch p-value: T.DIST.2T(|t|, df)", "p", "p"),
            (51, "Mann–Whitney: sum of ranks of the first half", "R1", "num2"),
            (52, "U = sum of ranks − n₁(n₁ + 1) ÷ 2", "U", "num2"),
            (53, "Expected U if the halves are alike = n₁n₂ ÷ 2", "EU", "num2"),
            (54, "z = (U − expected) ÷ √[n₁n₂(n₁ + n₂ + 1) ÷ 12]", "z", "num"),
            (55, "Mann–Whitney p-value", "pU", "p")]
    for r, label, key, fmt in rows:
        rep.row(r, label, fmt, C=res["C"][key], D=res["D"][key])
    rep.row(56, "Do the halves differ significantly?", "",
            **{c: yes(res[c]["p"] < ALPHA or res[c]["pU"] < ALPHA) for c in "CD"})

    # The account by half: Sharpe ratios (to the end of the account) and CAPM alphas (to French's last day)
    hday, T = regs["split_day"], len(da)
    lab_c = f"First half (to {text_dmy(da.date.iat[hday - 1])})"
    lab_d = f"Second half (from {text_dmy(da.date.iat[hday])})"
    rep.text("C60", lab_c)
    rep.text("D60", lab_d)
    rep.panel("4a. The account by half", {"C": lab_c, "D": lab_d})
    sh = lambda k, a, b: sharpe_daily(da["x_" + k].values[a:b]) * math.sqrt(DAYS_PER_YEAR)
    rep.row(61, "Sharpe ratio, Long + SPY, Window A", "num", C=sh("A", 0, hday), D=sh("A", hday, T))
    rep.row(62, "Sharpe ratio, Long + SPY, Window B", "num", C=sh("B", 0, hday), D=sh("B", hday, T))
    rep.row(63, "Sharpe ratio, B1 SPY buy and hold", "num", C=sh("B1", 0, hday), D=sh("B1", hday, T))
    for r, k in [(64, "A"), (68, "B")]:
        r1, r2 = regs[f"capm_{k}_h1"], regs[f"capm_{k}_h2"]
        p1, p2 = t_dist_2t(r1.t, r1.T - 2), t_dist_2t(r2.t, r2.T - 2)
        rep.row(r, f"CAPM alpha per year, Long + SPY account, Window {k}", "pct", C=r1.alpha_year, D=r2.alpha_year)
        rep.row(r + 1, "…t-statistic (Newey–West)", "num", C=r1.t, D=r2.t)
        rep.row(r + 2, "…p-value: T.DIST.2T(|t|, days − 2)", "p", C=p1, D=p2)
        rep.row(r + 3, f"Significant at the {text_pct(ALPHA)} level?", "", C=sig_word(p1, r1.alpha_year),
                D=sig_word(p2, r2.alpha_year))

    # 4b: year by year
    rep.panel("4b. Year by year", {"B": "Year", "C": "Trades", "D": "Mean A", "E": "Median A", "F": "Beat A",
                                   "G": "Mean B", "H": "Median B", "I": "Beat B", "J": "Label"})
    beat = lambda x: (x > 0).sum() / (x != 0).sum()
    for y in range(y0, y1 + 1):
        m = year == y
        r = 77 + (y - y0)
        rep.row(r, str(y), {"B": "year", "C": "int", "D": "pct", "E": "pct", "F": "pct", "G": "pct", "H": "pct",
                            "I": "pct"}, check_label=False,
                B=y, C=int(m.sum()), D=xA[m].mean(), E=float(np.median(xA[m])), F=beat(xA[m]),
                G=xB[m].mean(), H=float(np.median(xB[m])), I=beat(xB[m]), J=f"{y} (n={int(m.sum())})")
    rep.row(77 + (y1 - y0) + 1, "All years", {"C": "int", "D": "pct", "E": "pct", "F": "pct", "G": "pct", "H": "pct",
                                              "I": "pct"}, C=len(tr), D=xA.mean(), E=float(np.median(xA)),
            F=beat(xA), G=xB.mean(), H=float(np.median(xB)), I=beat(xB))

    # 4c: V1 quarterly rebalance additions only
    q = robust_block(rep, 138, "4c. V1: quarterly-rebalance additions only",
                     {"C": ("V1 quarterly, Window A", xA, Q, d), "D": ("V1 quarterly, Window B", xB, Q, d)},
                     {"C": vA, "D": vB})
    # 4d: the other variants
    rep.text("B164", f"V4 Without {EXCL_YEAR_1} and {EXCL_YEAR_2}")
    rep.panel("4d. V3 winsorising cut-offs", {"C": "Window A", "D": "Window B"})
    rep.row(170, f"Lower cut-off: {text_pct(WINSOR_LO)} percentile", "pct", C=loA, D=loB)
    rep.row(171, f"Upper cut-off: {text_pct(WINSOR_HI)} percentile", "pct", C=hiA, D=hiB)
    v4 = f"V4 without {EXCL_YEAR_1} and {EXCL_YEAR_2}"
    heads = {"C": "V2 off-cycle", "D": "V3 winsorised", "E": v4, "F": "V5 sell the day before", "G": "V6 trimmed"}
    rep.text("E174", v4)
    rep.text("E197", v4)
    vA_ = {c: vA for c in "CDEFG"}
    vB_ = {c: vB for c in "CDEFG"}
    dA = robust_block(rep, 175, "4d. Variants, Window A",
                      {"C": (heads["C"], xA, OFF, d), "D": (heads["D"], wA, ALL, d), "E": (heads["E"], xA, V4, d),
                       "F": (heads["F"], xV5, ALL, d), "G": (heads["G"], xA, V6A, d)}, vA_)
    dB = robust_block(rep, 198, "4d. Variants, Window B",
                      {"C": (heads["C"], xB, OFF, d), "D": (heads["D"], wB, ALL, d), "E": (heads["E"], xB, V4, d),
                       "F": (heads["F"], xV5, ALL, d), "G": (heads["G"], xB, V6B, d)}, vB_)

    # Putting it together
    same = lambda b, v: yes(verdict(b) == v)
    first = {"C": same(halves["C"], vA), "D": same(halves["E"], vB)}
    second = {"C": same(halves["D"], vA), "D": same(halves["F"], vB)}
    result = {c: ("Robust" if first[c] == second[c] == "Yes" else
                  "Partly robust" if "Yes" in (first[c], second[c]) else "Not robust") for c in "CD"}
    others = {"C": sum(verdict(b) == vA for b in dA.values()), "D": sum(verdict(b) == vB for b in dB.values())}
    rep.panel("Putting it together", {"C": "Window A", "D": "Window B"})
    rep.row(223, "First half gives Test 1's verdict?", "", **first)
    rep.row(224, "Second half gives Test 1's verdict?", "", **second)
    rep.row(225, "Result for this window", "", **result)
    rep.row(226, "For reference: V1 quarterly only, verdict", "", C=verdict(q["C"]), D=verdict(q["D"]))
    rep.row(227, f"For reference: other variants giving Test 1's verdict (of {len(dA)})", "int", **others)
    rep.panel("Answer", {"C": "Result"})
    overall = ("Robust" if result["C"] == result["D"] == "Robust" else
               "Not robust" if result["C"] == result["D"] == "Not robust" else "Partly robust")
    rep.row(11, "Test 4 result", "", C=overall)
    rep.panels.insert(0, rep.panels.pop())
    # the subsets whose average is significantly above / below zero (used by the Summary sheet)
    sig = []
    sig += [(h[c], sig_word(halves[c].pG, halves[c].meanG)) for c in "CDEF"]
    sig += [(f"V1 quarterly, Window {w}", sig_word(q[c].pG, q[c].meanG)) for c, w in [("C", "A"), ("D", "B")]]
    sig += [(heads[c] + ("" if c == "F" else ", Window A"), sig_word(dA[c].pG, dA[c].meanG)) for c in "CDEFG"]
    sig += [(heads[c] + ", Window B", sig_word(dB[c].pG, dB[c].meanG)) for c in "CDEG"]
    return rep, result, sig


# ---------------------------------------------------------------------------------------------------------------------
# 2.12 Summary
# ---------------------------------------------------------------------------------------------------------------------
def report_summary(t1, t2, t3, t4, sig):
    rep = Report("Summary", "the four tests in one place")
    rep.panel("Results by test", {"D": "Window A", "E": "Window B"})
    rep.row(12, "Test 1: per trade", "", D=t1["D"], E=t1["E"])
    rep.row(13, "Test 2: risk-adjusted", "", D=t2["D"], E=t2["E"])
    rep.row(14, "Test 3: as an account", "", D=t3["D"], E=t3["E"])
    rep.row(15, "Test 4: robustness", "", D=t4["C"], E=t4["D"])
    passes = {c: yes(t1[c] == t2[c] == t3[c] == "Works" and t4[k] == "Robust") for c, k in [("D", "C"), ("E", "D")]}
    rep.row(16, "Passes all four tests?", "", **passes)
    def lst(word):
        """The Test 4 subsets whose average is significant in that direction, as the sheet lists them."""
        names = [lab for lab, s in sig if s == word]
        return "; ".join(names) + "." if names else "None."

    rep.panel("Test 4 subsets with a significant average (clustered t-test)", {"C": "Subsets"})
    rep.row(20, "Significantly above zero", "", C=lst("Yes, above zero"))
    rep.row(21, "Significantly below zero", "", C=lst("Yes, below zero"))
    rep.panel("Verdict", {"D": "Decision"})
    rep.row(8, "Verdict", "", D="Invest" if "Yes" in passes.values() else "Do not invest")
    rep.panels.insert(0, rep.panels.pop())
    return rep


# ---------------------------------------------------------------------------------------------------------------------
# 2.13 Extra Checks sheet, E1-E5 (its date averages and ranks are in Extra Checks Workings)
# ---------------------------------------------------------------------------------------------------------------------
def report_extra(tr, da, regs):
    rep = Report("Extra Checks", "E1 growth by half, E2 buying at the announcement close, E3 from 2020, "
                                 "E4 break test, E5 quarterly additions")
    T = len(da)
    d, year, names = tr.announced.values, tr.year.values, tr.name.values
    xA, xB = tr.AR_A.values, tr.AR_B.values
    H1, H2 = year <= SPLIT_YEAR, year > SPLIT_YEAR
    y0, y1 = int(year.min()), int(year.max())
    keys = {"C": "A", "D": "B", "E": "B1"}

    # E1: growth of the accounts by half
    h = regs["split_day"]
    v = {c: da["value_" + k].values for c, k in keys.items()}
    g1 = {c: v[c][h - 1] ** (DAYS_PER_YEAR / h) - 1 for c in keys}
    g2 = {c: (v[c][-1] / v[c][h - 1]) ** (DAYS_PER_YEAR / (T - h)) - 1 for c in keys}
    cols = {"C": "Long + SPY, A", "D": "Long + SPY, B", "E": "B1 SPY"}
    rep.panel("E1. Growth by half (value of $1 from the Daily Account)", cols)
    split_date = to_date(da.date.iat[h - 1])
    rep.row(17, f"First half: last day ({text_dmy(split_date)})", "", **{c: split_date for c in keys})
    rep.row(18, "First half: trading days", "int", **{c: h for c in keys})
    rep.row(19, "First half: value of $1 at the end ($1 at the 31 Dec 2009 close)", "usd",
            **{c: v[c][h - 1] for c in keys})
    rep.row(20, f"First half: growth a year = value^({DAYS_PER_YEAR} ÷ days) − 1", "pct", **g1)
    rep.row(21, "Second half: trading days", "int", **{c: T - h for c in keys})
    rep.row(22, "Second half: value of $1 at the end", "usd", **{c: v[c][-1] for c in keys})
    rep.row(23, f"Second half: growth a year = (end ÷ start)^({DAYS_PER_YEAR} ÷ days) − 1", "pct", **g2)
    rep.row(24, "First half: growth a year above B1", "pct", C=g1["C"] - g1["E"], D=g1["D"] - g1["E"])
    rep.row(25, "Second half: growth a year above B1", "pct", C=g2["C"] - g2["E"], D=g2["D"] - g2["E"])
    rA, rB = regs["capm_A_h2"], regs["capm_B_h2"]
    rep.row(26, "Second half: CAPM alpha a year (Test 4a)", "pct", C=rA.alpha_year, D=rB.alpha_year)
    rep.row(27, "…p-value (Newey–West)", "p", C=t_dist_2t(rA.t, rA.T - 2), D=t_dist_2t(rB.t, rB.T - 2))

    # E2: bought at the announcement close (not tradable: an upper bound)
    full_rows = [(0, "N: trades", "N", "int"),
                 (1, "Mean abnormal return (each trade once)", "mean", "pct"),
                 (2, "Median", "median", "pct"),
                 (3, "% of trades that beat SPY (k ÷ non-zero trades)", "beat", "pct"),
                 (4, "Announcement dates (G)", "G", "int"),
                 (5, "Mean of the date averages (each date once)", "meanG", "pct"),
                 (6, "Clustered SE", "seG", "pct"),
                 (7, "Clustered t = mean of date averages ÷ SE", "tG", "num"),
                 (8, "p-value: T.DIST.2T(|t|, G − 1)", "pG", "p")]
    wil_rows = [(0, "Wilcoxon n (non-zero trades)", "n", "int"),
                (1, "W+: sum of ranks of positive trades", "W", "num2"),
                (2, "z = (W+ − n(n + 1) ÷ 4) ÷ √[n(n + 1)(2n + 1) ÷ 24]", "z", "num"),
                (3, "Wilcoxon p-value", "pW", "p")]
    e2 = {c: stat_block(tr[k].values, d)
          for c, k in [("C", "CAR_inc_open"), ("D", "CAR_inc_close"), ("E", "CAR_buy_open")]}
    rep.panel("E2. Bought at the announcement close (CAR since the announcement close; all trades)",
              {"C": "To inclusion open", "D": "To inclusion close", "E": "Overnight (Ann+1 open)"})
    for j, label, f, fmt in full_rows:
        rep.row(34 + j, label, fmt, **{c: getattr(b, f) for c, b in e2.items()})
    for j, label, f, fmt in wil_rows:
        rep.row(43 + j, label, fmt, **{c: getattr(b, f) for c, b in e2.items()})
    rep.row(47, "Sign test p-value: 2 × BINOM.DIST(min(k, n − k), n, 0.5)", "p", **{c: b.pS for c, b in e2.items()})
    hh = {"C": f"To the inclusion open, {y0}–{SPLIT_YEAR}", "D": f"To the inclusion open, {SPLIT_YEAR + 1}–{y1}",
          "E": f"To the inclusion close, {y0}–{SPLIT_YEAR}", "F": f"To the inclusion close, {SPLIT_YEAR + 1}–{y1}"}
    for c, lab in hh.items():
        rep.text(f"{c}50", lab)
    by_half = [("C", "CAR_inc_open", H1), ("D", "CAR_inc_open", H2),
               ("E", "CAR_inc_close", H1), ("F", "CAR_inc_close", H2)]
    e2h = {c: stat_block(tr[k].values[m], d[m]) for c, k, m in by_half}
    rep.panel("E2. By half", hh)
    for r, (j, label, f, fmt) in zip(range(51, 57), [full_rows[0]] + full_rows[4:9]):
        rep.row(r, label, fmt, **{c: getattr(b, f) for c, b in e2h.items()})

    # E3: from START_2020 onwards
    F20 = d >= np.datetime64(START_2020)
    F20XA, F20XB = F20 & (xA < xA[F20].max()), F20 & (xB < xB[F20].max())
    e3 = {c: stat_block(x[m], d[m]) for c, x, m in [("C", xA, F20), ("D", xA, F20XA), ("E", xB, F20), ("F", xB, F20XB)]}
    largest = lambda x, m, k: names[int(np.nonzero(x == np.sort(x[m])[::-1][k - 1])[0][0])]
    rep.panel(f"E3. Trades announced from {text_dmy(START_2020)}",
              {"C": "Window A", "D": "A without largest", "E": "Window B", "F": "B without largest"})
    e3_fields = [("N", "int"), ("mean", "pct"), ("median", "pct"), ("beat", "pct"), ("G", "int"), ("meanG", "pct"),
                 ("seG", "pct"), ("tG", "num"), ("pG", "p"), ("lo", "pct"), ("n", "int"), ("W", "num2"), ("z", "num"),
                 ("pW", "p"), ("mde", "pct")]
    e3_labels = ["N: trades",
                 "Mean abnormal return (each trade once)",
                 "Median",
                 "% of trades that beat SPY (k ÷ non-zero trades)",
                 "Announcement dates (G)",
                 "Mean of the date averages (each date once)",
                 "Clustered SE",
                 "Clustered t = mean of date averages ÷ SE",
                 "p-value: T.DIST.2T(|t|, G − 1)",
                 f"Lower end of the {text_pct(1 - ALPHA)} interval: mean − t* × SE",
                 "Wilcoxon n (non-zero trades)",
                 "W+: sum of ranks of positive trades",
                 "z = (W+ − n(n + 1) ÷ 4) ÷ √[n(n + 1)(2n + 1) ÷ 24]",
                 "Wilcoxon p-value",
                 f"MDE = ({ZCRIT:.2f} + {ZPOW:.2f}) × clustered SE"]
    for j, ((f, fmt), label) in enumerate(zip(e3_fields, e3_labels)):
        rep.row(63 + j, label, fmt, **{c: getattr(b, f) for c, b in e3.items()})
    rep.row(78, "Verdict (Test 1's rule)", "", **{c: verdict(b) for c, b in e3.items()})
    rep.row(79, "Largest trade left out", "", D=largest(xA, F20, 1), F=largest(xB, F20, 1))
    # the account from the first trading day of 2020
    p20 = regs["p20"]                                       # day number of the last day before START_2020
    n20 = T - p20
    g = {c: (v[c][-1] / v[c][p20 - 1]) ** (DAYS_PER_YEAR / n20) - 1 for c in keys}
    srd = {c: sharpe_daily(da["x_" + k].values[p20:]) for c, k in keys.items()}
    jk = {c: jobson_korkie(da["x_" + keys[c]].values[p20:], da.x_B1.values[p20:]) for c in "CD"}
    c20 = {c: regs[f"capm_{keys[c]}_2020"] for c in "CD"}
    rep.panel(f"E3. The account from {text_dmy(START_2020)}", cols)
    rep.row(83, "Start: value taken on this day's close", "", **{c: to_date(da.date.iat[p20 - 1]) for c in keys})
    rep.row(84, "Trading days from the start", "int", **{c: n20 for c in keys})
    rep.row(85, "Value of $1 at the start", "usd", **{c: v[c][p20 - 1] for c in keys})
    rep.row(86, "Value of $1 at the end", "usd", **{c: v[c][-1] for c in keys})
    rep.row(87, f"Growth a year = (end ÷ start)^({DAYS_PER_YEAR} ÷ days) − 1", "pct", **g)
    rep.row(88, "Growth a year above B1", "pct", C=g["C"] - g["E"], D=g["D"] - g["E"])
    rep.row(89, "Sharpe ratio, daily = average excess return ÷ its st. dev.", "num", **srd)
    rep.row(90, f"Sharpe ratio, annualised = daily × √{DAYS_PER_YEAR}", "num",
            **{c: s * math.sqrt(DAYS_PER_YEAR) for c, s in srd.items()})
    rep.row(91, "ρ: correlation with B1's daily excess return", "num", **{c: jk[c].rho for c in "CD"})
    rep.row(92, "Jobson–Korkie z (Memmel, 2003), as in Test 3", "num", **{c: jk[c].z for c in "CD"})
    rep.row(93, "…p-value (two-sided)", "p", **{c: jk[c].p for c in "CD"})
    rep.row(94, "CAPM alpha a year (daily α × days a year)", "pct", **{c: c20[c].alpha_year for c in "CD"})
    rep.row(95, "…t-statistic (Newey–West)", "num", **{c: c20[c].t for c in "CD"})
    rep.row(96, "…p-value: T.DIST.2T(|t|, days − 2)", "p", **{c: t_dist_2t(c20[c].t, c20[c].T - 2) for c in "CD"})

    # E4: one break in the mean of the date averages (sup-F, Bai and Perron critical values)
    ALL = np.ones(len(tr), bool)
    XA, XB = xA < xA.max(), xB < xB.max()                   # every trade except the largest in that window
    e4 = {c: sup_f_break(x[m], d[m]) for c, x, m in [("C", xA, ALL), ("D", xB, ALL), ("E", xA, XA), ("F", xB, XB)]}
    rep.panel("E4. Break test on the date averages (one value per announcement date, in date order)",
              {"C": "Window A", "D": "Window B", "E": "A without largest", "F": "B without largest"})
    e4_rows = [(104, "Announcement dates (G)", "G", "int"),
               (105, "Dates kept out at each end = INT(share × G)", "lo", "int"),
               (106, "sup-F: the largest F", "sup", "num"),
               (107, "Break after: the last date of the first period", "last", ""),
               (108, "…the first date of the second period", "next", ""),
               (109, "Dates before the break", "k", "int"),
               (110, "Dates after the break", "k2", "int"),
               (111, "Mean of the date averages before", "m1", "pct"),
               (112, "Mean of the date averages after", "m2", "pct"),
               (113, "Critical value at 5% (Settings)", "cv5", "num2"),
               (114, "Critical value at 1% (Settings)", "cv1", "num2"),
               (115, "Significant break?", "sig", "")]
    for r, label, f, fmt in e4_rows:
        rep.row(r, label, fmt, **{c: getattr(b, f) for c, b in e4.items()})
    rep.row(116, "Largest trade left out", "", E=largest(xA, ALL, 1), F=largest(xB, ALL, 1))

    # E5: quarterly additions (V1), by half and without the N_DROP_V1 largest winners
    Q = tr.quarterly.values == "Yes"
    drop = lambda x: Q & np.array([(x[Q] > xi).sum() >= N_DROP_V1 for xi in x])
    QXA, QXB = drop(xA), drop(xB)
    e5 = {c: stat_block(x[m], d[m]) for c, x, m in
          [("C", xA, Q), ("D", xA, Q & H1), ("E", xA, Q & H2), ("F", xA, QXA),
           ("G", xB, Q), ("H", xB, Q & H1), ("I", xB, Q & H2), ("J", xB, QXB)]}
    hq = {"C": "A: all", "D": f"A: {y0}–{SPLIT_YEAR}", "E": f"A: {SPLIT_YEAR + 1}–{y1}",
          "F": f"A: without {N_DROP_V1}", "G": "B: all", "H": f"B: {y0}–{SPLIT_YEAR}",
          "I": f"B: {SPLIT_YEAR + 1}–{y1}", "J": f"B: without {N_DROP_V1}"}
    for c in "DEHI":
        rep.text(f"{c}121", hq[c])
    rep.panel("E5. Quarterly-rebalance additions (V1)", hq)
    e5_rows = [(122, "N: trades", "N", "int"), (123, "Announcement dates (G)", "G", "int"),
               (124, "Mean abnormal return (each trade once)", "mean", "pct"),
               (125, "Mean of the date averages (each date once)", "meanG", "pct"), (126, "Clustered SE", "seG", "pct"),
               (127, "Clustered t = mean of date averages ÷ SE", "tG", "num"),
               (128, "p-value: T.DIST.2T(|t|, G − 1)", "pG", "p"),
               (129, f"MDE = ({ZCRIT:.2f} + {ZPOW:.2f}) × clustered SE", "mde", "pct")]
    for r, label, f, fmt in e5_rows:
        rep.row(r, label, fmt, **{c: getattr(b, f) for c, b in e5.items()})
    for k in (1, 2, 3):
        rep.row(129 + k, f"Left out: largest winner {k}", "",
                F=largest(xA, Q, k) if k <= N_DROP_V1 else "–", J=largest(xB, Q, k) if k <= N_DROP_V1 else "–")
    return rep


def sup_f_break(x, dates):
    """Sup-F test for one break in the mean of the date averages (Andrews 1993; Bai and Perron 1998).
    y = date averages in date order (G values). For k = LO..G-LO-1 dates in the first segment:
    SSR_k = sum of squares around each segment's mean, F_k = (SSR_0 - SSR_k) / (SSR_k / (G - 2))."""
    by_date = pd.Series(x).groupby(dates).mean().sort_index()
    y, G = by_date.values, len(by_date)
    lo = int(math.floor(BREAK_TRIM * G))
    cs = np.cumsum(y)
    q, SG = (y * y).sum(), cs[-1]
    ssr0 = q - SG ** 2 / G
    ks = np.arange(lo, G - lo)
    s = cs[ks - 1]
    ssr = q - s ** 2 / ks - (SG - s) ** 2 / (G - ks)
    F = (ssr0 - ssr) / (ssr / (G - 2))
    j = int(np.argmax(F))                                   # first k with the largest F
    k = int(ks[j])
    sup = float(F[j])
    sig = "Yes, at 1%" if sup > BAI_PERRON_CV_1 else "Yes, at 5%" if sup > BAI_PERRON_CV_5 else "No"
    return SimpleNamespace(G=G, lo=lo, sup=sup, last=to_date(by_date.index[k - 1]), next=to_date(by_date.index[k]),
                           k=k, k2=G - k, m1=s[j] / k, m2=(SG - s[j]) / (G - k), cv5=BAI_PERRON_CV_5,
                           cv1=BAI_PERRON_CV_1, sig=sig)

# ---------------------------------------------------------------------------------------------------------------------
# 2.14 Run every calculation
# ---------------------------------------------------------------------------------------------------------------------
def all_regressions(da):
    """The regressions of the Regression Workings sheet (day numbers are 1-based positions in the Daily Account)."""
    z = day_number(da.date, da.attrs["ff_last"])                         # last day with factors (4,190 in v12)
    h = day_number(da.date, dt.date(SPLIT_YEAR, 12, 31))                 # last day of the first half (2,013)
    p20 = day_number(da.date, START_2020 - dt.timedelta(days=1))         # last day before START_2020 (2,516)
    regs = {"full_T": z, "split_day": h, "p20": p20}
    for k in ["A", "B"]:
        y = "x_" + k
        regs[f"capm_{k}"] = regression(da, y, 1, z)                      # regressions 1, 2
        regs[f"ff4_{k}"] = regression(da, y, 1, z, four_factor=True)     # regressions 3, 4
        regs[f"capm_{k}_h1"] = regression(da, y, 1, h)                   # regressions 5, 7
        regs[f"capm_{k}_h2"] = regression(da, y, h + 1, z)               # regressions 6, 8
        regs[f"capm_{k}_2020"] = regression(da, y, p20 + 1, z)           # Extra Checks E3
    return regs


def run_calculations(path):
    raw = load_inputs(path)
    spy = spy_close_by_date(raw)
    beta = beta_estimation(raw)
    tr = trade_returns(raw, beta)
    hd = holding_days(raw, tr)
    da = daily_account(raw, hd, spy)
    regs = all_regressions(da)
    settings = report_settings(da)
    t1, t1a, t1b = report_test1(tr)
    t2 = report_test2(tr, beta, t1a, t1b)
    t3, t3res = report_test3(tr, hd, da, regs)
    t4, t4res, sig = report_test4(tr, da, regs, t1a, t1b)
    summary = report_summary({"D": t1.cells["D141"], "E": t1.cells["E141"]},
                             {"D": t2.cells["D122"], "E": t2.cells["E122"]}, t3res, t4res, sig)
    extra = report_extra(tr, da, regs)
    return SimpleNamespace(raw=raw, beta=beta, tr=tr, hd=hd, da=da, regs=regs,
                           reports=[settings, summary, t1, t2, t3, t4, extra],
                           workings=workings_reports(tr, beta, hd, da))


# =====================================================================================================================
# PART 3. CHECK (--check): compare every recomputed number with the workbook's cached values
# =====================================================================================================================
RESULT_WORDS = {"Works", "Doesn't work", "Mixed", "Inconclusive", "Robust", "Partly robust", "Not robust", "Yes", "No",
                "Yes, above zero", "Yes, below zero", "Yes, above 50%", "Yes, below 50%", "Invest", "Do not invest",
                "Yes, at 1%", "Yes, at 5%", "Not yet"}
REL_TOL, ABS_TOL = 1e-9, 1e-12


def _norm_text(s):
    return re.sub(r"\s+", " ", str(s)).strip()


def cell_rc(addr):
    """'D37' -> (37, 4)."""
    m = re.fullmatch(r"([A-Z]+)(\d+)", addr)
    col = 0
    for ch in m.group(1):
        col = col * 26 + ord(ch) - 64
    return int(m.group(2)), col


def cell_addr(r, c):
    """(37, 4) -> 'D37'."""
    letters = ""
    while c:
        c, rem = divmod(c - 1, 26)
        letters = chr(65 + rem) + letters
    return f"{letters}{r}"


def check_against_workbook(path, reports):
    """Cell by cell comparison with the values cached in the workbook (the last values Excel/LibreOffice computed).

    On each results sheet, every formula cell whose cached value is a number or a date must have been recomputed
    here and must agree: |a - b| <= 1e-9 * max(|a|, |b|) + 1e-12 for numbers (the small absolute floor is for
    values such as 4e-16 or 0), the same calendar day for dates. Text cells recomputed here (result words, the
    Summary lists, short labels and headers) must agree exactly. The long explanatory sentences are not compared.
    Before comparing, the label in column B of every row used must read as expected, so a moved row cannot pass by
    accident. On Settings, every constant of this script is compared with its cell; on the workings sheets (Trade
    Returns, Beta Estimation, Holding Days, Daily Account) every cell of every column recomputed here, after checking
    a few column headers. Returns True if all agree.
    """
    from openpyxl import load_workbook
    from openpyxl.worksheet.formula import ArrayFormula
    wb_f = load_workbook(path, read_only=True, data_only=False)          # formulas
    wb_v = load_workbook(path, read_only=True, data_only=True)           # cached values
    is_formula = lambda v: isinstance(v, ArrayFormula) or (isinstance(v, str) and v.startswith("="))

    def grid(ws):
        out = {}
        for r, row in enumerate(ws.iter_rows(min_row=1, min_col=1, values_only=True), start=1):
            for c, v in enumerate(row, start=1):
                if v is not None:
                    out[(r, c)] = v
        return out

    print("\n" + "=" * 140)
    print(f"CHECK  |  recomputed values against the values cached in {Path(path).name}")
    print("=" * 140)
    all_ok = True
    total = dict(num=0, num_ok=0, date=0, date_ok=0, text=0, text_ok=0, labels=0, labels_bad=0, missing=0, skipped=0)
    for rep in reports:
        if rep.sheet not in wb_v.sheetnames:
            if rep.compare == "listed" and rep.sheet != "Settings":
                print(f"\n{rep.sheet}: not in this workbook (a workbook straight from the data collection has no "
                      f"workings), so not compared.")
                continue
            print(f"\n{rep.sheet}: not in this workbook, so not compared"
                  + (" (the Extra Checks sheet first appears in v13)" if rep.sheet == "Extra Checks" else "") + ".")
            continue
        forms, vals = grid(wb_f[rep.sheet]), grid(wb_v[rep.sheet])
        st = dict(num=0, num_ok=0, date=0, date_ok=0, text=0, text_ok=0, labels=0, labels_bad=0, missing=0,
                  skipped=0)
        worst_rel, worst_abs = (0.0, "-"), (0.0, "-")
        problems, skipped, absent_rows = [], [], set()
        # 1. the label in column B of every row this script uses, and the headers of the workings columns
        for r, label in sorted(rep.labels.items()):
            got = vals.get((r, 2))
            if rep.sheet == "Settings" and got is None:
                absent_rows.add(r)                       # a constant added in a later version (v13: C23:C27)
                continue
            st["labels"] += 1
            if _norm_text(got) != _norm_text(label):
                st["labels_bad"] += 1
                problems.append(f"  LABEL   B{r}: workbook {got!r}, expected {label!r}")
        for addr, text in rep.headers.items():
            got = vals.get(cell_rc(addr))
            st["labels"] += 1
            if _norm_text(got) != _norm_text(text):
                st["labels_bad"] += 1
                problems.append(f"  HEADER  {addr}: workbook {got!r}, expected {text!r}")
        # 2. the cells: every formula cell of a results sheet, or every cell listed (Settings, workings sheets)
        if rep.compare == "listed":
            targets = sorted(cell_rc(a) for a in rep.cells if cell_rc(a)[0] not in absent_rows)
        else:
            targets = sorted(rc for rc, f in forms.items() if is_formula(f))
            in_sheet = set(targets)
            for a in rep.cells:
                if cell_rc(a) not in in_sheet:
                    problems.append(f"  NOTE    {a}: computed here, but not a formula cell in the workbook "
                                    f"(workbook value {vals.get(cell_rc(a))!r})")
        for (r, c) in targets:
            addr, wv = cell_addr(r, c), vals.get((r, c))
            mine = rep.cells.get(addr)
            blank = mine is None or (isinstance(mine, (float, np.floating)) and np.isnan(mine))
            if rep.compare == "listed" and (wv is None or isinstance(wv, str)) and blank:
                continue                                 # blank here, blank (or a note such as "too few days") there
            if wv is None:
                if rep.compare == "listed":
                    problems.append(f"  DIFF    {addr}: workbook blank, script {mine!r}")
                    st["num"] += 1
                continue
            label = _norm_text(vals.get((r, 2), ""))[:45] if rep.compare == "formulas" else rep.sheet
            if isinstance(wv, dt.datetime):
                st["date"] += 1
                if blank or isinstance(mine, str) or is_number(mine):
                    st["missing"] += mine is None
                    problems.append(f"  DIFF    {addr} [{label}]: workbook {wv.date()}, script {mine!r}")
                elif to_date(mine) == wv.date():
                    st["date_ok"] += 1
                else:
                    problems.append(f"  DIFF    {addr} [{label}]: workbook {wv.date()}, script {to_date(mine)}")
            elif is_number(wv):
                st["num"] += 1
                if not is_number(mine) or blank:
                    st["missing"] += mine is None
                    problems.append(f"  DIFF    {addr} [{label}]: workbook {wv!r}, script {mine!r}")
                    continue
                a, b = float(wv), float(mine)
                diff, size = abs(a - b), max(abs(a), abs(b))
                if diff <= REL_TOL * size + ABS_TOL:
                    st["num_ok"] += 1
                else:
                    problems.append(f"  DIFF    {addr} [{label}]: workbook {a!r}, script {b!r} "
                                    f"(relative difference {diff / size:.2e})")
                if size > 1e-10 and diff / size > worst_rel[0]:
                    worst_rel = (diff / size, addr)
                if diff > worst_abs[0]:
                    worst_abs = (diff, addr)
            elif isinstance(wv, str):
                if mine is not None:
                    st["text"] += 1
                    if str(mine) == wv:
                        st["text_ok"] += 1
                    else:
                        problems.append(f"  DIFF    {addr} [{label}]: workbook {wv!r}, script {mine!r}")
                elif c == 2 and r in rep.labels:
                    pass                                 # a row label written by a formula: checked in step 1
                elif wv in RESULT_WORDS:
                    st["text"] += 1
                    st["missing"] += 1
                    problems.append(f"  MISSING {addr} [{label}]: workbook {wv!r}, not computed here")
                else:
                    st["skipped"] += 1                   # an explanatory sentence
                    skipped.append(addr)
        bad = (st["num"] - st["num_ok"]) + (st["date"] - st["date_ok"]) + (st["text"] - st["text_ok"]) \
            + st["labels_bad"]
        all_ok &= bad == 0
        for k in st:
            total[k] += st[k]
        print(f"\n{rep.sheet}: {'all agree' if bad == 0 else 'DIFFERENCES FOUND'}")
        print(f"  numbers {st['num_ok']} of {st['num']} agree, dates {st['date_ok']} of {st['date']}, "
              f"text {st['text_ok']} of {st['text']}; labels {st['labels'] - st['labels_bad']} of {st['labels']}")
        if st["num"]:
            print(f"  largest relative difference {worst_rel[0]:.1e} ({worst_rel[1]}); "
                  f"largest absolute difference {worst_abs[0]:.1e} ({worst_abs[1]})")
        if absent_rows:
            print(f"  rows {min(absent_rows)}-{max(absent_rows)} are not in this version of the sheet (added in v13)")
        if skipped:
            print(f"  not compared, explanatory sentences ({len(skipped)}): {', '.join(skipped)}")
        for p in problems:
            print(p)
    print(f"\nTOTAL: numbers {total['num_ok']} of {total['num']} agree, dates {total['date_ok']} of {total['date']}, "
          f"text {total['text_ok']} of {total['text']}, labels {total['labels'] - total['labels_bad']} of "
          f"{total['labels']}; {total['missing']} not recomputed; {total['skipped']} sentences not compared.")
    if total["num"] == 0:
        print("RESULT: nothing to compare: the workbook holds no cached formula results "
              "(open it in Excel, let it recalculate and save it, then run --check again).")
        return False
    print("RESULT: every recomputed value matches the workbook." if all_ok else
          "RESULT: some values differ (listed above).")
    return all_ok


# =====================================================================================================================
# MAIN: the command line
# =====================================================================================================================
def main(argv=None):
    ap = argparse.ArgumentParser(description="S&P 500 index-inclusion strategy: recompute every statistic of the "
                                             "SPX Additions workbook from its raw data sheets.")
    ap.add_argument("workbook", nargs="?", help="path to 'SPX Additions v<N>.xlsx' (default: the newest one found "
                                                "next to this script, else in ../../data/prices)")
    ap.add_argument("--check", action="store_true", help="compare every recomputed number with the workbook's "
                                                         "cached values; exit code 1 if any differs")
    ap.add_argument("--rebuild-data", action="store_true", help="PART 1 first: re-collect the raw data from CRSP "
                                                                "(WRDS account), Yahoo Finance and Kenneth French's "
                                                                "library into a new workbook version, then analyse it")
    ap.add_argument("--rebuild-lists", action="store_true", help="with --rebuild-data: also rerun the helpers that "
                                                                 "built the candidate list of additions (CSV files)")
    ap.add_argument("--save-workings", metavar="FOLDER", help="also save Trade Returns, Holding Days and Daily "
                                                             "Account as CSV files in FOLDER")
    args = ap.parse_args(argv)

    path = original = find_workbook(args.workbook)
    if args.rebuild_data:
        if args.rebuild_lists:
            rebuild_addition_lists(path.parent)
        # The rebuilt workbook holds fresh raw data but no cached results (openpyxl does not calculate), so the
        # calculations run on it and --check compares them with the results cached in the original workbook.
        path = rebuild_event_data(path)

    print(f"SPX Additions analysis: reading the raw data sheets of {path}")
    res = run_calculations(path)
    print_workings(res.raw, res.tr, res.beta, res.hd, res.da)
    for rep in res.reports:
        print_report(rep)
    if args.save_workings:
        save_workings(args.save_workings, res.tr, res.beta, res.hd, res.da)
    if args.check:
        ok = check_against_workbook(original, res.reports + res.workings)
        return 0 if ok else 1
    return 0


if __name__ == "__main__":
    sys.exit(main())
